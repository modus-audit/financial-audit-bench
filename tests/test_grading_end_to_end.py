"""Workbook grading workflows, saved feedback, regrading, and formula recalculation.

Judge responses are deterministic fixtures. The LibreOffice workflow uses the real
installed executable and is skipped when LibreOffice is unavailable.
"""

from __future__ import annotations

import json
from contextlib import contextmanager
from pathlib import Path
from types import SimpleNamespace
import pytest
from openpyxl import Workbook, load_workbook
from financial_audit_bench.benchmark.grading import feedback, grade, judge_request
from financial_audit_bench.benchmark.grading import judge as judge_runtime
import shutil
from xml.etree import ElementTree
from zipfile import ZipFile
from financial_audit_bench.benchmark.grading import recalculate
from financial_audit_bench.benchmark.grading.recalculation_cache import RecalculationCache


@pytest.fixture
def judge_requests(monkeypatch):
    requests = []

    def complete(**kwargs):
        request = json.loads(kwargs['messages'][1]['content'])
        requests.append(request)
        assert kwargs['request_timeout_seconds'] == 120
        cells = request.get('candidate_answer_cells') or [request['cell']]
        passed = all(abs(float(c['value']) - 100) <= .01 for c in cells) if request['id'] == 'amount' else all(
            c['value'] == 'Confirmed' for c in cells)
        return SimpleNamespace(content=json.dumps({'checks': [{
            'id': request['id'], 'passed': passed, 'reason': 'Compared all candidate answers.',
            'evidence_cells': [c['id'] for c in cells],
        }]}), model='test-judge', usage={}, cost=.01, latency_seconds=.1)

    monkeypatch.setattr(judge_runtime, 'complete', complete)
    return requests


@pytest.fixture
def grading_case(tmp_path, monkeypatch, judge_requests):
    task, workspace = tmp_path / 'task', tmp_path / 'workspace'
    task.mkdir()
    workspace.mkdir()
    book = Workbook()
    sheet = book.active
    sheet.title = 'Cash'
    sheet['A1'] = 'January accounts'
    for column, value in enumerate(['Account', 'Amount', 'Conclusion'], 1):
        sheet.cell(3, column, value)
    sheet['A8'] = 'End accounts'
    template = task / 'template.xlsx'
    book.save(template)
    sheet['A4'], sheet['B4'], sheet['C4'] = 'Operating', 100, 'Confirmed'
    book.save(workspace / 'answer.xlsx')
    book.close()
    rubric = {
        'output': 'answer.xlsx',
        'comparison': {'numeric_absolute_tolerance': .01},
        'tabs': [{'id': 'cash', 'sheet': 'Cash', 'checks': [{
            'id': 'amount', 'kind': 'row_value', 'description': 'Confirmed amount',
            'key': {'Account': 'Operating'}, 'column': 'Amount', 'expected': 100,
            'comparison': 'numeric', 'end_header': ['End accounts'],
            'section': {'after': {'A': 'January accounts'}, 'before': {'A': 'End accounts'}},
        }, {
            'id': 'llm.conclusion', 'kind': 'llm', 'sheet': 'Cash',
            'description': 'Confirmation conclusion', 'expected': 'Confirmed',
            'extract': [{'sheet': 'Cash', 'selector': {
                'kind': 'key_exists', 'key': {'Account': 'Operating'},
                'end_header': ['End accounts']}, 'columns': ['Conclusion']}],
        }]}],
    }
    (task / 'rubric.json').write_text(json.dumps(rubric))

    @contextmanager
    def cached_values(paths):
        yield paths

    monkeypatch.setattr(grade, 'recalculated_workbooks', cached_values)
    return workspace, task, template, tmp_path / 'report.json'


def test_keyed_grading_and_sample_coverage_follow_edited_tables(grading_case, judge_requests):
    workspace, task, template, output = grading_case
    submission = workspace / 'answer.xlsx'
    keys = [{'Account': '0780', 'Check #': '41001'}, {'Account': '0780', 'Check #': '41002'}]
    for path in (template, submission):
        book = load_workbook(path)
        checks = book.create_sheet('Checks')
        checks.append(['Account', 'Check #', 'Amount', 'Test result', 'Description'])
        checks['A8'] = 'Coverage summary'
        if path == submission:
            checks.append(['Outside the table'])
            for row, key in enumerate(keys, 2):
                for column, value in enumerate([*key.values(), 100, 'Y', 'Tested check'], 1):
                    checks.cell(row, column, value)
            checks['E4'] = 'Conclusion: All eligible checks tested.'
            cash = book['Cash']
            cash.insert_rows(1, 2)
            cash.insert_rows(6)
            cash['A6'], cash['B6'] = 'Other account', 999
            cash['B7'] = 100.004  # Normal rounding, on the correct keyed row.
        book.save(path)
        book.close()
    rubric_path = task / 'rubric.json'
    rubric = json.loads(rubric_path.read_text())
    rubric['tabs'].append({'id': 'checks', 'sheet': 'Checks', 'checks': [{
        'id': 'checks.coverage', 'kind': 'selection_coverage', 'description': 'Test the January checks.',
        'key_columns': ['Account', 'Check #'], 'completion_columns': ['Amount', 'Test result'],
        'key_comparisons': {'Account': 'identifier', 'Check #': 'check_reference'},
        'eligible_keys': keys, 'mandatory_keys': keys, 'minimum_count': 2,
        'population_source': 'January statement: two eligible checks.', 'end_header': ['Coverage summary'],
        'summary_rows': {'column': 'Description', 'labels': ['Conclusion'],
                         'identity_columns': ['Account', 'Check #']},
    }]})
    rubric['tabs'][0]['checks'].append({
        'id': 'cash.notes', 'kind': 'notes_have_context', 'description': 'Link conclusions to records.',
        'note_columns': ['Conclusion'], 'primary_columns': ['Account', 'Amount'],
        'end_header': ['End accounts'],
    })
    rubric_path.write_text(json.dumps(rubric))
    before = submission.read_bytes()
    report = grade.evaluate(*grading_case)
    assert report['score'] == 1 and report['total_checks'] == 4
    assert judge_requests[0]['cell']['id'] == 'Cash!C7'
    assert submission.read_bytes() == before
    annotated = load_workbook(report['annotated_workbook'])
    assert 'Confirmed amount' in annotated['Cash']['B7'].comment.text
    assert annotated['Cash']['B6'].comment is None
    annotated.close()

    # A duplicate record and a narrative claim cannot replace the missing item.
    book = load_workbook(submission)
    book['Checks']['B3'] = '41001'
    book['Cash']['A7'] = 'Wrong account'
    book['Cash']['C8'] = 'A conclusion without any supporting record.'
    book.save(submission)
    book.close()
    failed = grade.evaluate(*grading_case)
    assert failed['score'] == 0
    assert len(judge_requests) == 1  # Missing keyed answers never reach the judge.
    assert all(c['passed'] is False for c in failed['checks'])


@pytest.mark.parametrize('contradictory', [False, True])
def test_numeric_duplicates_fail_while_semantic_records_are_judged(grading_case, judge_requests, contradictory):
    submission = grading_case[0] / 'answer.xlsx'
    book = load_workbook(submission)
    sheet = book['Cash']
    sheet['A5'], sheet['B5'], sheet['C5'] = 'Operating', 999 if contradictory else 100, 'Unresolved' if contradictory else 'Confirmed'
    sheet['A10'], sheet['B10'], sheet['C10'] = 'Operating', 999, 'Outside table'
    book.save(submission)
    book.close()
    report = grade.evaluate(*grading_case)
    assert report['score'] == (0 if contradictory else 0.5)
    assert {r['id'] for r in judge_requests} == {'llm.conclusion'}
    for request in judge_requests:
        assert 'table' in request
        assert {c['row'] for c in request['candidate_answer_cells']} == {4, 5}
    checks = {c['id']: c for c in report['checks']}
    assert checks['amount']['passed'] is False
    assert checks['amount']['error'] == 'ambiguous_lookup'
    assert checks['llm.conclusion']['passed'] is (not contradictory)
    assert report['workbook_highlighting']['status'] == 'complete'


@pytest.mark.parametrize('failure', ['provider', 'invalid_citation', 'trusted_template', 'missing_submission',
                                   'missing_value', 'missing_row', 'missing_header', 'missing_sheet'])
def test_grading_cli_distinguishes_failed_answers_from_evaluation_errors(grading_case, monkeypatch, failure):
    workspace, task, template, output = grading_case
    if failure in {'provider', 'invalid_citation'}:
        def complete(**kwargs):
            if failure == 'provider':
                raise TimeoutError('provider unavailable')
            return SimpleNamespace(content=json.dumps({'checks': [{
                'id': 'llm.conclusion', 'passed': True, 'reason': 'Unsupported citation',
                'evidence_cells': ['Cash!Z999'],
            }]}), model='test', usage={}, cost=.01, latency_seconds=.1)
        monkeypatch.setattr(judge_runtime, 'complete', complete)
    elif failure == 'trusted_template':
        book = load_workbook(template)
        book['Cash']['B3'] = 'Wrong column'
        book.save(template)
        book.close()
    elif failure == 'missing_submission':
        (workspace / 'answer.xlsx').unlink()
    else:
        book = load_workbook(workspace / 'answer.xlsx')
        sheet = book['Cash']
        if failure == 'missing_value':
            sheet['B4'] = None
        elif failure == 'missing_row':
            sheet['A4'] = 'Other account'
        elif failure == 'missing_header':
            sheet['A3'] = 'Unrecognized column'
        else:
            sheet.title = 'Other sheet'
        book.save(workspace / 'answer.xlsx')
        book.close()
    reward = output.with_name('reward.json')
    reward.write_text('{"reward": 1}')
    reward.with_suffix('.txt').write_text('1')
    monkeypatch.setattr('sys.argv', ['grade', '--workspace', str(workspace), '--task', str(task),
        '--template', str(template), '--output', str(output), '--reward'])
    missing = failure.startswith('missing_')
    expected_score = .5 if failure == 'missing_value' else 0
    if missing:
        grade.main()
        assert json.loads(reward.read_text()) == {'reward': expected_score}
    else:
        with pytest.raises(SystemExit, match='Evaluation failed'):
            grade.main()
        assert not reward.exists()
    report = json.loads(output.read_text())
    assert report['score'] == (expected_score if missing else None)
    if missing:
        assert all(c['passed'] is not None for c in report['checks'])
        assert next(c for c in report['checks'] if c['id'] == 'amount')['passed'] is False
    assert report['total_checks'] == 2
    assert not reward.with_suffix('.txt').exists()


def test_regrading_reuses_valid_judgments_and_refreshes_changed_inputs(grading_case, judge_requests, monkeypatch):
    workspace, task, template, output = grading_case
    first = grade.evaluate(*grading_case)
    cached = grade.evaluate(*grading_case)
    assert first['score'] == cached['score'] == 1
    assert len(judge_requests) == 1
    assert cached['semantic']['judge']['reused_requests'] == 1
    assert cached['semantic']['judge']['cost_usd'] == 0
    cache = next((output.parent / 'llm_judge').glob('request_*.json'))
    cache.write_text('{"fingerprint":')
    assert grade.evaluate(*grading_case)['score'] == 1
    assert len(judge_requests) == 2
    rubric_path = task / 'rubric.json'
    rubric = json.loads(rubric_path.read_text())
    rubric['tabs'][0]['checks'][1]['expected'] = 'The bank confirmation supports the balance.'
    rubric_path.write_text(json.dumps(rubric))
    assert grade.evaluate(*grading_case)['score'] == 1
    assert len(judge_requests) == 3
    book = load_workbook(workspace / 'answer.xlsx')
    book['Cash']['C4'] = 'Unresolved'
    book.save(workspace / 'answer.xlsx')
    book.close()
    assert grade.evaluate(*grading_case)['score'] == .5
    assert len(judge_requests) == 4
    monkeypatch.setattr(judge_request, 'PROMPT', judge_request.PROMPT + '\nRevised auditing instruction.')
    assert grade.evaluate(*grading_case)['score'] == .5
    assert len(judge_requests) == 5


def test_feedback_failure_keeps_reward_and_can_retry_without_grading(grading_case, monkeypatch):
    workspace, task, template, output = grading_case
    real_annotate = feedback.annotate_workbook

    def failed_feedback(*_args):
        assert json.loads(output.read_text())['score'] == 1
        raise OSError('feedback destination unavailable')

    monkeypatch.setattr(feedback, 'annotate_workbook', failed_feedback)
    monkeypatch.setattr('sys.argv', ['grade', '--workspace', str(workspace), '--task', str(task),
        '--template', str(template), '--output', str(output), '--reward'])
    grade.main()
    report = json.loads(output.read_text())
    assert report['workbook_highlighting']['status'] == 'error'
    assert json.loads(output.with_name('reward.json').read_text()) == {'reward': 1}

    def no_grading(*_args, **_kwargs):
        pytest.fail('Feedback recovery must reuse the saved verdicts')

    monkeypatch.setattr(feedback, 'annotate_workbook', real_annotate)
    monkeypatch.setattr(grade, 'recalculated_workbooks', no_grading)
    monkeypatch.setattr(judge_runtime, 'complete', no_grading)
    recovered = grade.regenerate_feedback(workspace, task, output)
    assert recovered['workbook_highlighting']['status'] == 'complete'
    assert recovered['checks'] == report['checks']
    assert recovered['workbook_highlighting']['summary']['score'] == report['score']
    assert Path(recovered['annotated_workbook']).is_file()
    changed = load_workbook(workspace / 'answer.xlsx')
    changed['Cash']['B4'] = 999
    changed.save(workspace / 'answer.xlsx')
    changed.close()
    with pytest.raises(ValueError, match='same submission'):
        grade.regenerate_feedback(workspace, task, output)


@pytest.mark.parametrize('automatic_only', [False, True])
def test_prefilled_template_never_earns_credit(tmp_path, monkeypatch, automatic_only, benchmark_dataset):
    from openpyxl import load_workbook
    task = benchmark_dataset / 'tasks/manufacturing_v1_01_accounts_payable'
    rubric = json.loads((task / 'tests/rubric.json').read_text())
    template = task / 'environment/workspace' / rubric['output']
    # Saving a formatting edit must not turn supplied cells into answers.
    workbook = load_workbook(template)
    workbook.active.column_dimensions['A'].width = 24
    workbook.save(tmp_path / rubric['output'])
    workbook.close()
    def unexpected(*args, **kwargs):
        raise AssertionError('Blank submissions need neither recalculation nor judge requests')
    monkeypatch.setattr(grade, 'recalculated_workbooks', unexpected)
    report = grade.evaluate(tmp_path, task / 'tests', template, tmp_path / 'report.json', automatic_only=automatic_only)
    assert report['score'] == 0
    assert all(c['passed'] is False for c in report['checks'])


@pytest.mark.parametrize('contradictory', [False, True])
def test_repeated_headers_and_field_label_notes_are_graded(grading_case, judge_requests, contradictory):
    workspace, task, template, output = grading_case
    submission = workspace / 'answer.xlsx'
    for path in (template, submission):
        book = load_workbook(path)
        sheet = book['Cash']
        sheet['A12'] = 'Effective interest rate (EIR)'
        if path == submission:
            sheet['B12'] = .06
            sheet['A13'] = 'Effective interest rate equals the stated rate because issuance costs are nil.'
            for col, value in enumerate(['Account', 'Amount', 'Conclusion'], 1):
                sheet.cell(5, col, value)
            sheet['A6'], sheet['B6'], sheet['C6'] = 'Operating', 999 if contradictory else 100, 'Unresolved' if contradictory else 'Confirmed'
            # A copied header outside the declared section cannot widen its scope.
            for col, value in enumerate(['Account', 'Amount', 'Conclusion'], 1):
                sheet.cell(15, col, value)
            sheet['A16'], sheet['B16'], sheet['C16'] = 'Operating', 999, 'Outside table'
        book.save(path)
        book.close()
    path = task / 'rubric.json'
    rubric = json.loads(path.read_text())
    rubric['tabs'][0]['checks'].extend([
        {'id': 'eir', 'kind': 'row_value', 'description': 'Annual rate',
         'match': {'key': {'A': 'Effective interest rate (EIR)'},
                   'comparisons': {'A': 'row_label'}}, 'column_letter': 'B',
         'expected': .06, 'comparison': 'numeric'},
        {'id': 'notes', 'kind': 'notes_have_context', 'description': 'Notes identify their records',
         'primary_columns': ['Account', 'Amount'], 'note_columns': ['Conclusion'],
         'end_header': ['End accounts'],
         'section': {'after': {'A': 'January accounts'}, 'before': {'A': 'End accounts'}}},
    ])
    path.write_text(json.dumps(rubric))
    before = submission.read_bytes()
    report = grade.evaluate(*grading_case)
    checks = {c['id']: c for c in report['checks']}
    assert checks['eir']['passed'] is True
    assert checks['eir']['matched_cells'] == ['B12']
    assert checks['notes']['passed'] is True
    assert checks['amount']['passed'] is False
    assert checks['llm.conclusion']['passed'] is (not contradictory)
    assert report['score'] == (0.5 if contradictory else 0.75)
    assert submission.read_bytes() == before
    for request in judge_requests:
        assert {c['row'] for c in request['candidate_answer_cells']} == {4, 6}
        assert set(request['table']['candidate_rows']) == {4, 6}


def test_shared_recalculation_produces_the_same_grade_without_worker_calc(grading_case, monkeypatch):
    from financial_audit_bench.benchmark.grading import recalculate
    from financial_audit_bench.benchmark.grading.recalculation_cache import RecalculationCache
    workspace, task, template, output = grading_case
    original = workspace / 'answer.xlsx'
    before = original.read_bytes()
    single = grade.evaluate(*grading_case)

    @contextmanager
    def convert(sources, **kwargs):
        yield sources
    monkeypatch.setattr(recalculate, 'recalculated_workbooks', convert)
    cache = RecalculationCache(output.parent / 'host-cache')
    assert cache.prepare([original]) == {}
    monkeypatch.setattr(grade, 'recalculated_workbooks', lambda *_: pytest.fail('worker launched Calc'))
    batch = grade.evaluate(*grading_case, recalculation_cache=cache.directory)
    assert single['checks'] == batch['checks']
    assert single['score'] == batch['score'] == 1
    assert batch['semantic']['judge']['reused_requests'] == 1
    assert original.read_bytes() == before
    cache.resolve(original).write_bytes(b'tampered')
    with pytest.raises(recalculate.RecalculationError, match='missing or changed'):
        grade.evaluate(*grading_case, recalculation_cache=cache.directory)


def test_automatic_only_uses_single_file_and_still_validates_semantic_rules(grading_case, judge_requests):
    from financial_audit_bench.benchmark.grading.rubric import load_rubric

    workspace, task, template, output = grading_case
    assert sorted(p.name for p in task.glob('*.json')) == ['rubric.json']
    report = grade.evaluate(workspace, task, template, output, automatic_only=True)
    assert report['score'] == 1 and report['total_checks'] == 1
    assert judge_requests == []
    assert set(report['provenance']['rubrics_sha256']) == {'rubric.json'}
    assert 'manual_review' not in report
    # Selecting a mode must not hide defects in the trusted answer key.
    path = task / 'rubric.json'
    rubric = json.loads(path.read_text())
    rubric['tabs'][0]['checks'][1]['expected'] = ''
    path.write_text(json.dumps(rubric))
    with pytest.raises(ValueError, match='expected answer'):
        load_rubric(task, automatic_only=True)


def test_an_unchanged_correct_review_retains_credit(grading_case):
    workspace, task, template, output = grading_case
    # A review starts from a completed submission; grade against the blank template.
    correct = grade.evaluate(workspace, task, template, output, automatic_only=True)
    assert correct["score"] == 1
    incorrect = grade.evaluate(workspace, task, workspace / "answer.xlsx",
                               output.with_name("incorrect.json"), automatic_only=True)
    assert incorrect["score"] == 0


@pytest.mark.skipif(
    shutil.which("soffice") is None,
    reason="Requires local LibreOffice (soffice on PATH)",
)
def test_batched_recalculation_replaces_stale_dependent_caches(tmp_path: Path, monkeypatch) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append([3, "=A1*2", "=B1+1"])
    uncached = tmp_path / "uncached.xlsx"
    workbook.save(uncached)
    workbook.close()

    # Model a workbook whose precedent changed from 1 to 3 without refreshing
    # either formula cache. Ordinary LibreOffice conversion retains 2 and 3.
    source = tmp_path / "stale.xlsx"
    ns = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    with ZipFile(uncached) as original, ZipFile(source, "w") as stale:
        for member in original.infolist():
            data = original.read(member)
            if member.filename == "xl/worksheets/sheet1.xml":
                xml = ElementTree.fromstring(data)
                for cell, value in (("B1", "2"), ("C1", "3")):
                    xml.find(f".//s:c[@r='{cell}']/s:v", ns).text = value
                data = ElementTree.tostring(xml)
            stale.writestr(member, data)
    before = source.read_bytes()

    commands = []
    popen = recalculate.subprocess.Popen
    def launch(command, **kwargs):
        commands.append(command)
        return popen(command, **kwargs)
    monkeypatch.setattr(recalculate.subprocess, "Popen", launch)
    cache = RecalculationCache(tmp_path / "batch-cache")
    assert cache.prepare([source, uncached]) == {}
    assert len(commands) == 1
    assert sum(str(arg).endswith(".xlsx") for arg in commands[0]) == 2
    for original in (source, uncached):
        values = load_workbook(cache.resolve(original), data_only=True)
        formulas = load_workbook(cache.resolve(original), data_only=False)
        assert [values.active[cell].value for cell in ("A1", "B1", "C1")] == [3, 6, 7]
        assert formulas.active["B1"].value == "=A1*2"
        assert formulas.active["C1"].value == "=B1+1"
        values.close()
        formulas.close()

    assert cache.prepare([source, uncached]) == {}
    assert len(commands) == 1
    assert source.read_bytes() == before
