"""Accounts-payable aging and subsequent-disbursement workbooks."""

from __future__ import annotations

from decimal import Decimal
from pathlib import Path
from typing import Any

from financial_audit_bench.synthetic_binders.priors.graph_construction.projection.foundational_support import (
    _add_schedule_sheet,
)
from financial_audit_bench.synthetic_binders.priors.graph_construction.projection.materialize import (
    MONEY_FORMAT,
    save_workbook,
    open_pbc_sheet,
    write_header,
)
from financial_audit_bench.synthetic_binders.priors.graph_construction.projection.materialize_identity import (
    client_id_map,
)

World = dict[str, Any]


def render_combined_ap_aging(path: Path, world: World, title: str):
    """AP-to-GL reconciliation, aging detail, and PO cross-reference."""
    from financial_audit_bench.synthetic_binders.priors.graph_construction.projection.identity import (
        package_identity,
    )

    aging = world.get("accounts_payable_aging") or []
    wb, rec = open_pbc_sheet("AP-to-GL Reconciliation", "AP-01", world)
    rec.title = "AP-to-GL Reconciliation"
    cell_map = []
    # The world carries entity/component rows followed by the reconciled account total.
    # Only the complete reconciliation belongs on this tab; emitting every component
    # produced blank GL cells plus a redundant total.
    reconciliation_rows = [
        row
        for row in aging
        if row.get("gl_balance") is not None and row.get("difference") is not None
    ]
    write_header(
        rec,
        ["Account", "Subledger", "GL Balance (Credit)", "Difference", "Conclusion"],
    )
    for row in reconciliation_rows:
        difference = Decimal(str(row["difference"]))
        rec.append(
            [
                row["payable_gl_account_id"],
                row["open_amount"],
                row["gl_balance"],
                None,
                "Reconciled" if difference == 0 else "Investigate difference",
            ]
        )
        reconciliation_row = rec.max_row
        rec.cell(reconciliation_row, 4).value = (
            # The subledger is displayed as a positive liability while the GL export
            # retains its signed credit. Adding the two is the genuine reconciliation
            # difference; subtracting would double the payable balance.
            f"=B{reconciliation_row}+C{reconciliation_row}"
        )
        rec.cell(reconciliation_row, 4).number_format = MONEY_FORMAT
        for column, (field, value) in enumerate(
            (
                ("Account", row["payable_gl_account_id"]),
                ("Subledger total", row["open_amount"]),
                ("GL total", row["gl_balance"]),
                ("Differences", row["difference"]),
            ),
            start=1,
        ):
            if value is None:
                continue
            from openpyxl.utils import get_column_letter

            cell_map.append(
                {
                    "field": field,
                    "cell": f"{get_column_letter(column)}{rec.max_row}",
                    "value": str(value),
                }
            )
    open_items = world.get("accounts_payable_open_item") or []
    invoices = {
        row["vendor_invoice_id"]: row for row in world.get("vendor_invoice") or []
    }
    vendors = {row["vendor_id"]: row for row in world.get("vendor") or []}
    receipts = {
        row["goods_receipt_id"]: row for row in world.get("goods_receipt") or []
    }
    receipt_numbers = {
        row["goods_receipt_id"]: f"Receipt {index:05d}"
        for index, row in enumerate(
            sorted(
                world.get("goods_receipt") or [],
                key=lambda value: (
                    value.get("receipt_date", ""),
                    value["goods_receipt_id"],
                ),
            ),
            start=1,
        )
    }
    identity = package_identity(world)
    client_ids = client_id_map(world)
    detail = _add_schedule_sheet(wb, "AP Aging Detail", world)
    write_header(
        detail,
        [
            "Line",
            "Vendor",
            "Invoice",
            "Invoice Date",
            "Due Date",
            "Original Amount",
            "Open Balance",
            "Aging Bucket",
            "GL Account",
            "PO Number",
            "Support Status",
        ],
    )
    detail_first_row = detail.max_row + 1
    po_rows = []
    for line, item in enumerate(open_items, start=1):
        invoice = invoices[item["vendor_invoice_id"]]
        receipt_id = invoice.get("goods_receipt_id")
        receipt = receipts.get(receipt_id)
        raw_po_id = (receipt or {}).get("purchase_order_id") or ""
        po_id = client_ids.get(raw_po_id, raw_po_id)
        status = "Vouched" if receipt else "Not vouched"
        number = identity.invoice_numbers.get(
            item["vendor_invoice_id"], invoice.get("invoice_reference", "")
        )
        detail.append(
            [
                line,
                vendors[invoice["vendor_id"]]["name"],
                number,
                invoice["document_date"],
                item["due_date"],
                item["original_amount"],
                item["open_amount"],
                item["aging_bucket"],
                item["payable_gl_account_id"],
                po_id,
                status,
            ]
        )
        if po_id:
            po_rows.append(
                [
                    line,
                    vendors[invoice["vendor_id"]]["name"],
                    number,
                    po_id,
                    receipt_numbers.get(receipt_id, ""),
                    status,
                ]
            )
    detail_last_row = detail.max_row
    open_total = sum(
        (Decimal(str(item["open_amount"])) for item in open_items),
        Decimal("0.00"),
    )
    if open_items:
        detail.append(
            [
                "Total",
                len(open_items),
                "",
                "",
                "",
                f"=SUM(F{detail_first_row}:F{detail_last_row})",
                f"=SUM(G{detail_first_row}:G{detail_last_row})",
            ]
        )
        for column in (6, 7):
            detail.cell(detail.max_row, column).number_format = MONEY_FORMAT
        reconciliation_open_total = sum(
            (Decimal(str(row["open_amount"])) for row in reconciliation_rows),
            Decimal("0.00"),
        )
        if open_total != reconciliation_open_total:
            raise ValueError(
                "AP aging detail does not reconcile to the AP-to-GL control: "
                f"detail={open_total}; reconciliation={reconciliation_open_total}"
            )
    po = _add_schedule_sheet(wb, "PO Cross-Reference", world)
    write_header(
        po,
        [
            "AP Line",
            "Vendor",
            "Invoice",
            "PO Number",
            "Receipt Number",
            "Support Status",
        ],
    )
    for row in po_rows:
        po.append(row)
    save_workbook(wb, path)
    return cell_map


def render_subsequent_disbursements(
    path: Path, world: World, title: str
) -> list[dict[str, Any]]:
    """Render AP-04 with a complete bank-debit classification bridge."""

    from financial_audit_bench.synthetic_binders.priors.graph_construction.projection.documents.subsequent_period_statement import (
        build_subsequent_statement_activity,
    )
    from financial_audit_bench.synthetic_binders.priors.graph_construction.projection.op_ap_subsequent import (
        subsequent_open_invoice_settlements,
    )

    settlements = subsequent_open_invoice_settlements(world)
    activity = build_subsequent_statement_activity(world)
    if activity is None:
        raise ValueError("AP-04 requires the subsequent-period bank statement")
    client_ids = client_id_map(world)

    wb, payments = open_pbc_sheet(title, "AP-04", world)
    payments.title = "Subsequent Disbursements"
    write_header(
        payments,
        [
            "Payment ID",
            "Payment Date",
            "Bank Clearing Date",
            "Method",
            "Payment Reference",
            "Payee",
            "Invoice",
            "Invoice Date",
            "Receipt / Service Date",
            "GL Account",
            "Amount",
            "Source Bank Account",
        ],
    )
    cell_map: list[dict[str, Any]] = []
    for settlement in settlements:
        payments.append(
            [
                settlement["payment_id"],
                settlement["payment_date"],
                settlement["clearing_date"],
                settlement["payment_method"],
                settlement["payment_reference"],
                settlement["payee"],
                settlement["invoice_number"],
                settlement["document_date"],
                settlement["service_period_end"],
                settlement["payable_gl_account_id"],
                settlement["amount"],
                settlement["source_bank_account"],
            ]
        )
        cell_map.extend(
            (
                {
                    "field": "Payment ID",
                    "cell": f"A{payments.max_row}",
                    "value": settlement["payment_id"],
                },
                {
                    "field": "Amount",
                    "cell": f"K{payments.max_row}",
                    "value": settlement["amount"],
                },
            )
        )
    listing_total = sum(
        (Decimal(str(row["amount"])) for row in settlements), Decimal("0.00")
    )
    payments.append(
        ["Total", len(settlements), "", "", "", "", "", "", "", "", listing_total]
    )

    debits = [row for row in activity["rows"] if Decimal(str(row["amount"])) < 0]
    unclassified = [
        row
        for row in debits
        if not row.get("classification")
        or not row.get("source_family")
        or not row.get("source_record_ids")
    ]
    if unclassified:
        raise ValueError(
            "January statement contains unclassified debits: "
            + ", ".join(str(row.get("description")) for row in unclassified)
        )

    electronic = [row for row in debits if not row["is_check"]]
    checks = [row for row in debits if row["is_check"]]
    classified_count = len(electronic) + len(checks)
    classified_total = sum(
        (abs(Decimal(str(row["amount"]))) for row in debits), Decimal("0.00")
    )
    ap_rows = [row for row in debits if row.get("ap_population_id")]
    ap_count = len(ap_rows)
    ap_total = sum(
        (abs(Decimal(str(row["amount"]))) for row in ap_rows), Decimal("0.00")
    )
    if ap_count != len(settlements) or ap_total != listing_total:
        raise ValueError(
            "AP-04 does not reconcile to the January statement: "
            f"listing {len(settlements)} / {listing_total}; "
            f"statement {ap_count} / {ap_total}"
        )

    control = _add_schedule_sheet(wb, "Bank Debit Control", world)
    write_header(control, ["Source control", "Count", "Amount"])
    electronic_total = sum(
        (abs(Decimal(str(row["amount"]))) for row in electronic), Decimal("0.00")
    )
    check_total = sum(
        (abs(Decimal(str(row["amount"]))) for row in checks), Decimal("0.00")
    )
    for label, count, amount in (
        ("Statement — electronic withdrawals", len(electronic), electronic_total),
        ("Statement — checks paid", len(checks), check_total),
        ("Statement — all debits", len(debits), classified_total),
        ("Classified statement debits", classified_count, classified_total),
        ("Unclassified difference", len(debits) - classified_count, Decimal("0.00")),
        ("AP-04 listing", len(settlements), listing_total),
        ("AP-04 rows matched to statement", ap_count, ap_total),
        (
            "AP-04 reconciliation difference",
            len(settlements) - ap_count,
            listing_total - ap_total,
        ),
    ):
        control.append([label, count, amount])

    control.append([])
    write_header(
        control,
        [
            "Statement Line ID",
            "Section",
            "Bank Date",
            "Description",
            "Amount",
            "Classification",
            "Source Family",
            "Source Record ID(s)",
            "AP-04 Payment ID",
            "Disposition",
        ],
    )
    for index, row in enumerate(debits, start=1):
        ap_id = row.get("ap_population_id") or ""
        disposition = (
            "Included in AP-04 subsequent-disbursement population"
            if ap_id
            else f"Excluded from AP-04 — resolved in {row['source_family']} cycle"
        )
        control.append(
            [
                f"JAN{activity['year']}-DEBIT-{index:04d}",
                "Checks Paid" if row["is_check"] else "Electronic Withdrawals",
                row["date"],
                row["description"],
                abs(Decimal(str(row["amount"]))),
                row["classification"],
                str(row["source_family"]).replace("_", " ").title(),
                ", ".join(
                    str(client_ids.get(value, value))
                    for value in row["source_record_ids"]
                ),
                ap_id,
                disposition,
            ]
        )
    control.append(["Total classified", len(debits), "", "", classified_total])

    from financial_audit_bench.synthetic_binders.priors.graph_construction.projection.ap import (
        append_outstanding_check_ap_relief,
    )

    cell_map.extend(append_outstanding_check_ap_relief(wb, world))

    save_workbook(wb, path)
    return cell_map
