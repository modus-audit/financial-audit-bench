"""Correctness-only finalization for generated client workbooks."""

from __future__ import annotations

import re
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from financial_audit_bench.synthetic_binders.priors.graph_construction.projection.materialize_identity import (
    client_id_map,
)
from financial_audit_bench.synthetic_binders.priors.graph_construction.projection.public_presentation import (
    present_public_text,
)

_IDENTIFIER = re.compile(r"\b[A-Z][A-Z0-9]*(?:-[A-Z0-9_]+)+\b")


def finalize_workbooks(file_map: list[dict[str, Any]], world: dict[str, Any]) -> None:
    """Flatten mapped formulas and remove private graph identifiers."""
    public_ids = client_id_map(world)
    public_id_fields = {
        "public_journal_entry_id",
        "journal_entry_line_id",
        "parent_journal_entry_line_id",
    }
    internal_ids = {
        value
        for population in world.values()
        for row in (population if isinstance(population, list) else [population])
        if isinstance(row, dict)
        for field, value in row.items()
        if isinstance(value, str)
        and (
            field.endswith("_id") or field in {"invoice_reference", "payment_reference"}
        )
        and "account_id" not in field
        and field not in public_id_fields
    }

    def public_text(value: str) -> str:
        def replace(match: re.Match[str]) -> str:
            token = match.group(0)
            if token not in internal_ids:
                return token
            return str(
                public_ids.get(token)
                or f"Record reference {' '.join(token.split('-')[1:])}"
            )

        return _IDENTIFIER.sub(replace, present_public_text(value))

    for entry in file_map:
        path = Path(entry["path"])
        if path.suffix.lower() != ".xlsx" or not path.is_file():
            continue
        workbook = load_workbook(path)
        default_sheet = workbook.active.title
        ties = {
            (str(row.get("sheet") or default_sheet), str(row["cell"])): row
            for row in entry.get("cell_map") or []
            if row.get("cell")
        }
        for sheet in workbook.worksheets:
            for row in sheet.iter_rows():
                for cell in row:
                    tie = ties.get((sheet.title, cell.coordinate))
                    if cell.data_type == "f" and tie is not None:
                        cell.value = tie["value"]
                    if isinstance(cell.value, str) and not cell.value.startswith("="):
                        rendered = public_text(cell.value)
                        if rendered != cell.value:
                            cell.value = rendered
                            if tie is not None:
                                tie["value"] = rendered
        workbook.save(path)
        workbook.close()
