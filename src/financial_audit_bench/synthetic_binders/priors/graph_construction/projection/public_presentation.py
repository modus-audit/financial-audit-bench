"""Last-mile presentation rules for public workbook renderers."""

from __future__ import annotations

import re
from pathlib import Path
from typing import Any


_SNAKE_TOKEN = re.compile(
    r"(?<![A-Za-z0-9_])([a-z0-9]+(?:_[a-z0-9]+)+)(?![A-Za-z0-9_])"
)
_GL_ACCOUNT_TOKEN = re.compile(r"\bGL-([A-Z0-9_]+(?:-[A-Z0-9_]+)*)\b")
_LINEAGE_TOKEN = re.compile(
    r"\b(?:JOURNAL|SUPPORT|VOUCHER|VENDOR|AP-PAYMENT|PREPAID-ITEM|PREPAID-SUPPORT|"
    r"WIP-CONTRACT|EMP|USER|NOISE|ITEM|RECON|BOOK|DEBIT|CREDIT|DEBT|ACCRUAL|STK|INVOICE|REF|RNI|GROUP|CR|BANK|AR|CUST|CI|CIL|CUT|FUL|CC)-[A-Za-z0-9][A-Za-z0-9-]*"
    r"|\b[A-Z0-9-]+-(?:ERROR|CORRECTION)\b"
)
_NUMBERED_INTERNAL_TOKEN = re.compile(
    r"\b([A-Z][A-Z0-9]*(?:-[A-Z0-9_]+)*)-(0\d{2,3})\b"
)


def _lineage_label(token: str) -> str:
    """Turn a graph identifier into a neutral client-system reference."""
    parts = token.split("-")
    head = parts[0]
    vocabulary = {
        "ACCRUAL": "Accrual",
        "APPROVER": "Approver",
        "AUDIT": "Audit",
        "BEN": "Benefits",
        "HIRE": "Hire",
        "PAYMENT": "Payment",
        "PAYROLL": "Payroll",
        "RNI": "RNI",
    }
    tail = " ".join(
        vocabulary.get(part, part.replace("_", " ").title()) for part in parts[1:]
    )
    labels = {
        "JOURNAL": "Journal reference",
        "SUPPORT": "Support reference",
        "VOUCHER": "Voucher reference",
        "VENDOR": "Vendor reference",
        "EMP": "Employee record",
        "USER": "User role",
        "NOISE": "System entry",
        "ITEM": "Reconciling item",
        "RECON": "Reconciliation reference",
        "BOOK": "Book reference",
        "ACCRUAL": "Accrual reference",
        "DEBIT": "Debit reference",
        "CREDIT": "Credit reference",
        "DEBT": "Debt reference",
        "STK": "Stock record",
        "INVOICE": "Invoice reference",
        "REF": "System reference",
        "RNI": "Received-not-invoiced reference",
        "GROUP": "Coverage group",
        "CR": "Cash receipt",
        "BANK": "Bank account",
        "AR": "Confirmation reference",
        "CUST": "Customer account",
        "CI": "Customer invoice",
        "CIL": "Invoice line",
        "CUT": "Cutoff selection",
        "FUL": "Fulfillment record",
        "CC": "Customer contract",
    }
    if token.startswith("PREPAID-ITEM-"):
        return f"Prepaid item {' '.join(parts[2:])}"
    if token.startswith("PREPAID-SUPPORT-"):
        return f"Prepaid support {' '.join(parts[2:])}"
    if token.startswith("WIP-CONTRACT-"):
        return f"Contract reference {' '.join(parts[2:])}"
    if token.startswith("AP-PAYMENT-"):
        return f"Payment reference {' '.join(parts[2:])}"
    return f"{labels.get(head, 'Record reference')} {tail}".strip()


def present_public_identifiers(value: str) -> str:
    """Humanize internal identifiers without changing surrounding vocabulary."""
    value = _GL_ACCOUNT_TOKEN.sub(
        lambda match: f"Source account {' '.join(match.group(1).split('-'))}", value
    )
    value = _LINEAGE_TOKEN.sub(lambda match: _lineage_label(match.group(0)), value)
    return _NUMBERED_INTERNAL_TOKEN.sub(
        lambda match: (
            match.group(0)
            if match.group(1).startswith("JEL-")
            else (
                f"{' '.join(match.group(1).split('-')).title()} reference "
                f"{int(match.group(2)):03d}"
            )
        ),
        value,
    )


def present_public_text(value: str) -> str:
    """Humanize standalone enum tokens and internal record identifiers."""
    value = present_public_identifiers(value)
    return _SNAKE_TOKEN.sub(
        lambda match: match.group(1).replace("_", " ").capitalize(), value
    )


def apply_workbook_presentation_boundary(
    file_map: list[dict[str, Any]], run_dir: Path
) -> None:
    """Apply the public vocabulary boundary to completed XLSX artifacts."""
    from openpyxl import load_workbook

    for artifact in file_map:
        if artifact.get("request_id") == "DOC-PACKAGE-INDEX":
            continue
        raw_path = Path(str(artifact.get("path") or ""))
        path = raw_path if raw_path.is_absolute() else run_dir / raw_path
        if path.suffix.lower() != ".xlsx" or not path.is_file():
            continue
        workbook = load_workbook(path)
        changed = False
        changed_cells: set[tuple[str, str]] = set()
        for sheet in workbook.worksheets:
            for row in sheet.iter_rows():
                for cell in row:
                    value = cell.value
                    if not isinstance(value, str) or value.startswith("="):
                        continue
                    rendered = present_public_text(value)
                    if rendered != value:
                        cell.value = rendered
                        changed = True
                        changed_cells.add((sheet.title, cell.coordinate))
        if not changed:
            continue
        for tie in artifact.get("cell_map") or []:
            address = tie.get("cell")
            if not address:
                continue
            sheet_name = tie.get("sheet") or workbook.active.title
            if (sheet_name, address) in changed_cells:
                tie["value"] = str(workbook[sheet_name][address].value or "")
        workbook.save(path)
