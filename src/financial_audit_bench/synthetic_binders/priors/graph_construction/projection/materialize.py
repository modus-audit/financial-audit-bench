"""Materialize Core PBC artifacts as client-style workbooks."""

from __future__ import annotations

import re
from decimal import Decimal
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from financial_audit_bench.synthetic_binders.data_catalog import load_registry
from financial_audit_bench.synthetic_binders.priors.graph_construction.projection.materialize_identity import (
    client_id_map,
)
from financial_audit_bench.synthetic_binders.priors.graph_construction.projection.request_resolution import (
    NAMED_EVIDENCE_PROVIDERS,
)

_REQUEST_VALUES = load_registry("registry.pbc.requests.v1").values
INDIRECT_REQUESTS = set(NAMED_EVIDENCE_PROVIDERS)
REQUEST_SATISFIED_BY_DOCUMENTS = dict(_REQUEST_VALUES["request_document_kinds"])
SPECIALIZED_WORKBOOK_REQUESTS = frozenset(
    {"AP-01", "AP-03", "AP-04", "AP-06", "AP-07", "AR-01", "AR-07", "INV-10"}
)
MANUFACTURING_SPECIALIZED_WORKBOOK_REQUESTS = frozenset({"REV-01", "REV-02", "REV-03"})
HEADER_KIND_BY_REQUEST = {
    request_id: "point_in_time"
    for request_id in _REQUEST_VALUES["point_in_time_requests"]
}


# Values that omit a schedule instead of rendering placeholder rows.
OMITTED_VALUES = {"not_applicable", None, "", "Not required", "not required"}
# Workbook titles use client-facing language, not auditor request text.


# -- workbook styling ---------------------------------------------------------
# Shared workbook styles.
TITLE_FONT = Font(bold=True, size=14, color="FF1F3864")
HEADER_FONT = Font(bold=True, size=10, color="FF1F3864")
HEADER_FILL = PatternFill("solid", fgColor="FFD9E1F2")
HEADER_BORDER = Border(bottom=Side(style="medium", color="FF8EA9DB"))
TOTAL_FONT = Font(bold=True)
TOTAL_BORDER = Border(top=Side(style="thin", color="FF404040"))
# Total/subtotal rows get a rule and bold. "net income"/"net loss" are totals;
# a bare "net " would wrongly catch line items like "Net investment income".
_TOTAL_PREFIXES = (
    "total",
    "subtotal",
    "grand total",
    "balance forward",
    "net income",
    "net loss",
)


def write_header(ws, cells: list) -> None:
    """Append a table column-header row and give it the shaded header band."""
    ws.append(cells)
    for cell in ws[ws.max_row]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.border = HEADER_BORDER


def filename(request_id: str, request_name: str) -> str:
    """Return a readable filename derived from the authoritative request."""
    name = "".join(c for c in request_name if c.isalnum() or c in " -&").strip()
    return name or request_id


# Non-tabular families are routed to text documents or specialized workbook renderers
# below. The shared renderer is intentionally limited to the remaining keyed schedules
# and scalar control sheets.
def materialize_pbc_workbooks(
    projection: dict[str, list[dict[str, Any]]],
    output_dir: Path,
    world: dict[str, Any],
) -> list[dict[str, Any]]:
    """Write one workbook per applicable request; return the file map."""
    lineage_by_request: dict[str, list[dict[str, Any]]] = {}
    for row in projection["core_pbc_field_lineage"]:
        lineage_by_request.setdefault(row["request_id"], []).append(row)

    file_map = []
    for artifact in projection["core_pbc_artifacts"]:
        if artifact["applicability_status"] == "out_of_scope":
            continue
        request_id = artifact["request_id"]
        if request_id == "AP-01" and not world.get("accounts_payable_aging"):
            # A payroll/accrual-enabled entity can legitimately have no trade AP
            # population. In that case neither an AP aging nor its embedded AP-to-GL
            # reconciliation exists as a client report.
            artifact["evidence_status"] = "no_reportable_activity"
            continue
        if request_id in INDIRECT_REQUESTS:
            if request_id == "AP-18" and not (
                world.get("ap_payment")
                or any(
                    movement.get("transaction_class") == "check"
                    and not movement.get("ap_payment_id")
                    for movement in world.get("cash_movement") or []
                )
            ):
                # AP-18 is satisfied by the complete journal export. When the AP module
                # and residual-check population are both empty, there is no journal to
                # deliver and no SURL listing whose completeness needs to be
                # established.
                artifact["evidence_status"] = "no_reportable_activity"
            continue
        if request_id in REQUEST_SATISFIED_BY_DOCUMENTS:
            # The documents channel owns these outputs; do not create field-table copies.
            from financial_audit_bench.synthetic_binders.priors.graph_construction.projection.documents import (
                document_is_applicable,
            )

            doc_kind = REQUEST_SATISFIED_BY_DOCUMENTS[request_id]
            if not document_is_applicable(doc_kind, world):
                artifact["evidence_status"] = "no_reportable_activity"
            continue
        name = filename(request_id, artifact["request_name"])
        path = output_dir / artifact["family"] / f"{name}.xlsx"
        path.parent.mkdir(parents=True, exist_ok=True)
        client_title = artifact["request_name"]
        business_type = str(world.get("business_type") or "")
        manufacturing_revenue = bool(
            request_id in {"REV-01", "REV-02", "REV-03"}
            and business_type == "manufacturing"
        )
        staffing_revenue_recon = bool(
            request_id == "REV-02" and business_type == "staffing_services"
        )
        if (
            request_id in SPECIALIZED_WORKBOOK_REQUESTS
            or request_id in MANUFACTURING_SPECIALIZED_WORKBOOK_REQUESTS
        ) and (
            request_id not in {"REV-01", "REV-02", "REV-03"}
            or manufacturing_revenue
            or staffing_revenue_recon
        ):
            # Render specialized schedules instead of the generic field table.
            from financial_audit_bench.synthetic_binders.priors.graph_construction.projection.foundational import (
                render_accrual_rollforward,
                render_subsequent_disbursements,
                render_received_not_invoiced,
                render_inventory_cutoff_population,
                render_combined_ap_aging,
                render_combined_ar_aging,
                render_allowance_workbook,
                render_manufacturing_revenue_detail,
                render_manufacturing_revenue_reconciliation,
                render_manufacturing_revenue_policy,
            )
            from financial_audit_bench.synthetic_binders.priors.graph_construction.projection.ap import (
                render_vendor_invoice_support,
            )

            renderer = {
                "AP-03": render_vendor_invoice_support,
                "AP-04": render_subsequent_disbursements,
                "AP-06": render_received_not_invoiced,
                "AP-07": render_accrual_rollforward,
                "INV-10": render_inventory_cutoff_population,
                "AR-01": render_combined_ar_aging,
                "AR-07": render_allowance_workbook,
                "AP-01": render_combined_ap_aging,
                "REV-01": render_manufacturing_revenue_detail,
                "REV-02": render_manufacturing_revenue_reconciliation,
                "REV-03": render_manufacturing_revenue_policy,
            }[request_id]
            cell_map = renderer(path, world, client_title)
        else:
            cell_map = _write_workbook(
                path,
                artifact,
                lineage_by_request.get(request_id, []),
                world,
                client_title,
            )
        if cell_map is None:
            artifact["evidence_status"] = "no_reportable_activity"
            continue
        file_map.append(
            {
                "request_id": request_id,
                "family": artifact["family"],
                "path": str(path),
                "cell_map": cell_map,
            }
        )
    return file_map


def sheet_title(title: str) -> str:
    """Sanitize a content title into a valid xlsx sheet name."""
    cleaned = "".join(c for c in title if c not in "[]:*?/\\").strip()
    if len(cleaned) > 31:
        cleaned = cleaned[:31]
        if " " in cleaned:
            cleaned = cleaned.rsplit(" ", 1)[0]
    return cleaned or "Schedule"


def open_pbc_sheet(title: str, request_id: str, world=None):
    """Create one target-package workbook with a compact client masthead."""
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title(title)
    if world:
        ws.append([world["company_context"]["legal_name"]])
        ws.append([title])
        year = int(world["fiscal_calendar"]["fiscal_year"])
        point_in_time = HEADER_KIND_BY_REQUEST.get(request_id) == "point_in_time"
        ws.append(
            [
                (
                    f"As of December 31, {year}"
                    if point_in_time
                    else f"For the year ended December 31, {year}"
                )
            ]
        )
    else:
        ws.append([title])
    ws.append([])
    return wb, ws


def append_tie(ws, field: str, value: Any) -> dict[str, Any]:
    """Emit a scalar tie row; the cell/value shape is the validation contract."""
    ws.append([field, str(value)])
    return {"field": field, "cell": f"B{ws.max_row}", "value": str(value)}


def fixed_asset_support_display(value: Any) -> Any:
    """Render capitalization-policy lineage as a client-facing source label."""
    if not isinstance(value, str) or "CAP-POLICY-" not in value:
        return value

    def policy_label(match: re.Match[str]) -> str:
        asset_class = match.group(1).replace("_", " ").replace("-", " ").title()
        return f"{asset_class} capitalization policy"

    return re.sub(r"CAP-POLICY-([A-Z0-9_-]+)", policy_label, value)


def _write_workbook(
    path: Path,
    artifact: dict[str, Any],
    lineage: list[dict[str, Any]],
    world: dict[str, Any],
    client_title: str | None = None,
) -> list[dict[str, Any]] | None:
    """Render projected scalars and equal-length columns into one workbook."""
    wb, ws = open_pbc_sheet(
        client_title or artifact["request_name"],
        artifact["request_id"],
        world,
    )
    header_overrides = (
        {"Quantity": "Hours", "Price": "Bill rate"}
        if artifact["request_id"] == "AR-03"
        and world.get("business_type") == "staffing_services"
        else {}
    )

    def _label(field: str) -> str | None:
        return header_overrides.get(field, field)

    cell_map: list[dict[str, Any]] = []
    scalars: list[tuple[str, Any]] = []
    dict_tables: list[tuple[str, list[dict[str, Any]]]] = []
    column_groups: dict[int, list[tuple[str, list[Any]]]] = {}
    for row in lineage:
        value = row["generated_value"]
        field = row["expected_field"]
        if isinstance(value, list) and value and isinstance(value[0], dict):
            dict_tables.append((field, value))
        elif isinstance(value, list):
            items = [item for item in value if item not in OMITTED_VALUES]
            if not items:
                continue
            column_groups.setdefault(len(items), []).append((field, items))
        elif isinstance(value, dict):
            dict_tables.append(
                (field, [{"item": k, "value": v} for k, v in value.items()])
            )
        elif isinstance(value, bool):
            continue
        elif value not in OMITTED_VALUES:
            scalars.append((field, value))

    if not scalars and not column_groups and not dict_tables:
        return None

    user_names = {
        row["accounting_user_id"]: row["user_name"]
        for row in world.get("accounting_user", [])
    }
    client_ids = client_id_map(world)

    def display(value: Any) -> Any:
        if isinstance(value, str):
            value = client_ids.get(value, user_names.get(value, value))
        return present(value)

    for field, value in scalars:
        label = _label(field)
        if label is None:
            continue
        rendered = display(value)
        ws.append([label, rendered])
        cell_map.append({"field": field, "cell": f"B{ws.max_row}", "value": rendered})
    if scalars:
        ws.append([])

    for length, fields in column_groups.items():
        write_header(ws, [_label(field) or field for field, _ in fields])
        for index in range(length):
            rendered = [display(items[index]) for _, items in fields]
            ws.append(rendered)
            for col, ((field, _), value) in enumerate(zip(fields, rendered), 1):
                cell_map.append(
                    {
                        "field": field,
                        "cell": f"{get_column_letter(col)}{ws.max_row}",
                        "value": value,
                    }
                )
        ws.append([])

    for field, rows in dict_tables:
        ws.append([_label(field) or field])
        columns = list(rows[0])
        write_header(ws, [column.replace("_", " ").capitalize() for column in columns])
        for item in rows:
            rendered = [display(item.get(column)) for column in columns]
            ws.append(rendered)
            for col, (column, rendered_value) in enumerate(zip(columns, rendered), 1):
                cell_map.append(
                    {
                        "field": f"{field}.{column}",
                        "cell": f"{get_column_letter(col)}{ws.max_row}",
                        "value": rendered_value,
                    }
                )
        ws.append([])

    save_workbook(wb, path)
    return cell_map


_SNAKE = re.compile(r"^[a-z0-9]+(_[a-z0-9]+)+$")
_LONG_DECIMAL = re.compile(r"^-?\d+\.\d{3,}$")
# Any plain numeric string (no thousands separators; leading zeros stay text — ZIPs and
# padded references are identifiers): typed as a number at the presentation boundary.
_PLAIN_NUMBER = re.compile(r"^-?(0|[1-9]\d*)(\.\d+)?$")
_ISO_DATE = re.compile(r"^\d{4}-\d{2}-\d{2}$")

MONEY_FORMAT = "#,##0.00;(#,##0.00)"
DATE_FORMAT = "mm/dd/yyyy"
PERCENT_FORMAT = "0.00%"
_RATE = re.compile(r"^-?0\.\d{2,}$")
_RATE_LABEL = re.compile(r"rate|yield|escalation|apr|%", re.IGNORECASE)
_TWO_DECIMAL = re.compile(r"^-?\d+\.\d{2}$")


def _money_value(value: Any) -> float | None:
    """Parse a client-rendered money string to a number, or None if it isn't one."""
    if isinstance(value, (int, float, Decimal)) and not isinstance(value, bool):
        return float(value)
    t = str(value).strip()
    negative = t.startswith("(") and t.endswith(")")
    t = t.strip("()").replace("$", "").replace(",", "").strip()
    if negative:
        t = f"-{t}"
    if not _TWO_DECIMAL.match(t):
        return None
    try:
        return float(t)
    except ValueError:
        return None


def finalize_workbook(wb) -> None:
    """Type and style every sheet so it reads like a real export/schedule."""
    from datetime import date as _date

    for ws in wb.worksheets:
        for row in ws.iter_rows():
            label = row[0].value if row else None
            for cell in row:
                value = cell.value
                if isinstance(value, str):
                    money = _money_value(value)
                    if (
                        _RATE.match(value)
                        and isinstance(label, str)
                        and _RATE_LABEL.search(label)
                    ):
                        cell.value = float(value)
                        cell.number_format = PERCENT_FORMAT
                    elif money is not None:
                        cell.value = money
                        cell.number_format = MONEY_FORMAT
                    elif _ISO_DATE.match(value):
                        try:
                            cell.value = _date.fromisoformat(value)
                        except ValueError:
                            continue
                        cell.number_format = DATE_FORMAT
                elif isinstance(value, float) and not isinstance(value, bool):
                    cell.number_format = MONEY_FORMAT
        # Second pass: a bare-integer string ('15000') inside a column that carries
        # typed money cells is a money amount a hand-rolled renderer appended as text —
        # type it. Leading-zero strings stay text (ZIPs, padded references), and columns
        # without money cells (quantities, check numbers) are untouched.
        money_columns = {
            cell.column
            for row in ws.iter_rows()
            for cell in row
            if cell.number_format == MONEY_FORMAT
        }
        for row in ws.iter_rows():
            for cell in row:
                value = cell.value
                if (
                    cell.column in money_columns
                    and isinstance(value, str)
                    and _PLAIN_NUMBER.match(value)
                ):
                    cell.value = float(value)
                    cell.number_format = MONEY_FORMAT
        _style_layout(ws)


def save_workbook(wb, path: Path) -> None:
    """Finalize and save one generated workbook."""
    path.parent.mkdir(parents=True, exist_ok=True)
    finalize_workbook(wb)
    wb.save(path)


def _style_layout(ws) -> None:
    """Apply a small, readable schedule style."""
    rows = list(ws.iter_rows())
    populated_rows = [
        row for row in rows if any(c.value not in (None, "") for c in row)
    ]
    if not populated_rows:
        return
    first = next(c for c in populated_rows[0] if c.value not in (None, ""))
    first.font = TITLE_FONT

    header_row = next(
        (
            row[0].row
            for row in rows
            if any(cell.fill and cell.fill.patternType for cell in row)
        ),
        None,
    )
    if header_row is None:
        for row in rows[2:]:
            cells = [cell for cell in row if cell.value not in (None, "")]
            if len(cells) >= 2 and all(isinstance(cell.value, str) for cell in cells):
                header_row = cells[0].row
                for cell in cells:
                    cell.font = HEADER_FONT
                    cell.fill = HEADER_FILL
                    cell.border = HEADER_BORDER
                break

    widths: dict[int, int] = {}
    for row in rows:
        cells = [cell for cell in row if cell.value not in (None, "")]
        if (
            cells
            and isinstance(cells[0].value, str)
            and cells[0].value.casefold().startswith(_TOTAL_PREFIXES)
        ):
            for cell in cells:
                cell.font = TOTAL_FONT
                cell.border = TOTAL_BORDER
        for cell in cells:
            widths[cell.column] = max(widths.get(cell.column, 0), len(str(cell.value)))
    for column, width in widths.items():
        ws.column_dimensions[get_column_letter(column)].width = min(
            max(width + 2, 10), 52
        )
    if header_row is not None:
        ws.freeze_panes = f"A{header_row + 1}"
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape" if len(widths) >= 7 else "portrait"


def present(value: Any) -> Any:
    """Client-facing presentation boundary."""
    if value is None:
        return ""
    if isinstance(value, (list, tuple)):
        return ", ".join(str(present(item)) for item in value)
    if isinstance(value, bool):
        return "Yes" if value else "No"
    if isinstance(value, str):
        if _SNAKE.match(value):
            return value.replace("_", " ").capitalize()
        if _PLAIN_NUMBER.match(value):
            # Preserve numeric types and precision for ratios and rates below one.
            number = float(value)
            if _LONG_DECIMAL.match(value) and abs(number) >= 1:
                return round(number, 2)
            if "." in value:
                return number
            return int(value)
        return value
    return value
