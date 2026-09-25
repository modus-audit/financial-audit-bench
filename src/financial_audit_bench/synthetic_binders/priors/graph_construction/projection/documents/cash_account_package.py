"""Compact client cash listing and year-end reconciliation workbook."""

from __future__ import annotations

from decimal import Decimal
from pathlib import Path
from typing import Any

from openpyxl import Workbook

from financial_audit_bench.synthetic_binders.priors.graph_construction.projection.documents._shared import (
    World,
)
from financial_audit_bench.synthetic_binders.priors.graph_construction.projection.documents.bank_reconciliation_controls import (
    calculate_bank_reconciliation,
    require_resolved_reconciliation,
    require_year_end_bank_statement,
)
from financial_audit_bench.synthetic_binders.priors.graph_construction.projection.materialize import (
    finalize_workbook,
    write_header,
)
from financial_audit_bench.synthetic_binders.priors.graph_construction.identities import (
    cash_gl_account_id,
)


def _accounts(world: World, year_end: dict[str, Any]) -> list[dict[str, Any]]:
    operating = world["bank_account"]
    cash_gl_id = cash_gl_account_id(world["general_ledger_account"])
    trial_balance = {
        row["gl_account_id"]: Decimal(str(row["closing_balance"]))
        for row in world["trial_balance_account"]
    }
    gl_names = {
        row["gl_account_id"]: f"{row['account_number']} - {row['name']}"
        for row in world["general_ledger_account"]
    }
    accounts = [
        {
            "id": operating["bank_account_id"],
            "institution": operating["bank_name"],
            "name": operating["account_name"],
            "number": operating["masked_account_number"],
            "currency": operating["currency_code"],
            "restriction": operating["restriction_status"].replace("_", " ").title(),
            "gl": gl_names[cash_gl_id],
            "bank": Decimal(str(year_end["ending_balance"])),
            "book": trial_balance[cash_gl_id],
        }
    ]
    for row in world.get("payroll_bank_account") or []:
        balance = Decimal(str(row["closing_balance"]))
        accounts.append(
            {
                "id": row["bank_account_id"],
                "institution": row["bank_name"],
                "name": row["account_name"],
                "number": row["masked_account_number"],
                "currency": world["company_context"]["currency_code"],
                "restriction": "Unrestricted",
                "gl": "Payroll clearing",
                "bank": balance,
                "book": balance,
            }
        )
    return accounts


def render(
    world: World,
    output_dir: Path,
) -> list[dict[str, Any]]:
    statements = world.get("bank_statement_month") or []
    if not statements:
        return []
    period_end = str(world["fiscal_calendar"]["end_date"])
    operating_id = world["bank_account"]["bank_account_id"]
    year_end = require_year_end_bank_statement(
        statements,
        bank_account_id=operating_id,
        statement_period_end=period_end,
    )
    accounts = _accounts(world, year_end)

    workbook = Workbook()
    listing = workbook.active
    listing.title = "Cash Accounts"
    listing.append([world["company_context"]["legal_name"]])
    listing.append([f"Cash Account Listing as of {period_end}"])
    listing.append([])
    write_header(
        listing,
        [
            "Institution",
            "Account Name",
            "Account Number",
            "Currency",
            "Restriction",
            "GL Account",
            "Bank Balance",
            "Book Balance",
        ],
    )
    for account in accounts:
        listing.append(
            [
                account["institution"],
                account["name"],
                account["number"],
                account["currency"],
                account["restriction"],
                account["gl"],
                account["bank"],
                account["book"],
            ]
        )

    detail = workbook.create_sheet("Reconciling Items")
    detail.append([world["company_context"]["legal_name"]])
    detail.append([f"Year-End Reconciling Items as of {period_end}"])
    detail.append([])
    write_header(
        detail,
        [
            "Account",
            "Item ID",
            "Type",
            "Book Date",
            "Bank Date",
            "Amount",
            "Subsequent Clearing Date",
            "Status",
            "Management Explanation",
        ],
    )

    for account in accounts:
        reconciliation = workbook.create_sheet((account["name"] + " Rec")[:31])
        reconciliation.append([world["company_context"]["legal_name"]])
        reconciliation.append([f"{account['name']} Reconciliation as of {period_end}"])
        reconciliation.append([])
        controls = calculate_bank_reconciliation(
            bank_account_id=account["id"],
            statement_period_end=period_end,
            statement_balance=account["bank"],
            book_balance=account["book"],
            reconciling_items=world.get("bank_reconciling_item") or [],
        )
        require_resolved_reconciliation(controls, bank_account_id=account["id"])
        write_header(reconciliation, ["Reconciliation", "Amount"])
        for label, amount in (
            ("Bank statement balance", controls.statement_balance),
            ("Add: deposits in transit", controls.deposits_in_transit),
            ("Less: outstanding checks and charges", controls.outstanding_checks),
            ("Adjusted bank balance", controls.adjusted_bank_balance),
            ("Book balance", controls.book_balance),
            ("Bank items not yet recorded", controls.bank_items_not_in_books),
            ("Adjusted book balance", controls.adjusted_book_balance),
            ("Difference", controls.difference),
        ):
            reconciliation.append([label, amount])
        for item in controls.all_items:
            explanation = item.get("management_explanation")
            if not explanation and item["item_kind"] == "outstanding_check":
                explanation = "Issued before year end and cleared after year end"
            elif not explanation and item["item_kind"] == "deposit_in_transit":
                explanation = "Recorded before year end and credited after year end"
            detail.append(
                [
                    account["name"],
                    item["bank_reconciling_item_id"],
                    item["item_kind"].replace("_", " ").title(),
                    item.get("book_date"),
                    item.get("bank_date"),
                    item["amount"],
                    item.get("subsequent_clearing_date"),
                    item["status"].replace("_", " ").title(),
                    explanation or "Pending resolution",
                ]
            )

    finalize_workbook(workbook)
    path = output_dir / "cash" / "Cash Account Listing and Reconciliations.xlsx"
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    return [
        {
            "request_id": "DOC-CASH-ACCOUNT-PACKAGE",
            "family": "cash",
            "path": str(path),
            "cell_map": [],
        }
    ]
