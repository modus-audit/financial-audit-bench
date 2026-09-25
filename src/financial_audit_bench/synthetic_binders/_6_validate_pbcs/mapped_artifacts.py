"""Tie delivered text and workbook cells to recorded lineage maps."""

from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


def _cell_matches(actual: Any, expected: Any) -> bool:
    import datetime

    if actual is None:
        actual = ""
    if actual == expected:
        return True
    if isinstance(actual, (datetime.date, datetime.datetime)):
        return actual.isoformat().replace("T00:00:00", "") == str(expected).replace(
            "T00:00:00", ""
        )

    def number(value: Any) -> float:
        if isinstance(value, str):
            text = value.strip()
            negative = text.startswith("(") and text.endswith(")")
            text = text.strip("()").replace("$", "").replace(",", "").strip()
            return float(f"-{text}" if negative else text)
        return float(value)

    try:
        return number(actual) == number(expected)
    except (TypeError, ValueError):
        return False


def _validate_line_map(path: Path, entry: dict) -> tuple[int, list[dict]]:
    """Tie a text deliverable to its recorded line lineage."""
    tied = 0
    failures: list[dict] = []
    lines = path.read_text().splitlines()
    for mapped in entry["line_map"]:
        line = (
            lines[mapped["line"] - 1].rstrip() if mapped["line"] <= len(lines) else ""
        )
        if line.endswith(str(mapped["value"])):
            tied += 1
        else:
            failures.append(
                {
                    "request_id": entry["request_id"],
                    "field": mapped["field"],
                    "line": mapped["line"],
                    "expected": mapped["value"],
                    "actual": line,
                }
            )
    return tied, failures


def _validate_cell_map(path: Path, entry: dict) -> tuple[int, list[dict]]:
    """Tie an XLSX deliverable to its recorded cell lineage in one pass."""
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        cells_by_sheet: dict[str, dict[str, Any]] = {}
        for sheet_name in {
            str(mapped.get("sheet") or workbook.sheetnames[0])
            for mapped in entry["cell_map"]
        }:
            if sheet_name not in workbook.sheetnames:
                cells_by_sheet[sheet_name] = {}
                continue
            sheet = workbook[sheet_name]
            cells_by_sheet[sheet_name] = {
                f"{get_column_letter(column)}{row_index}": cell.value
                for row_index, row in enumerate(sheet.iter_rows(), 1)
                for column, cell in enumerate(row, 1)
                if cell.value is not None
            }
    finally:
        workbook.close()

    failures: list[dict] = []
    tied = 0
    for mapped in entry["cell_map"]:
        sheet_name = str(mapped.get("sheet") or workbook.sheetnames[0])
        actual = cells_by_sheet.get(sheet_name, {}).get(mapped["cell"])
        if _cell_matches(actual, mapped["value"]):
            tied += 1
        else:
            # openpyxl returns typed dates for date-formatted cells. Failure diagnostics
            # are persisted in the validation and admissibility JSON contracts, so
            # retain the value while making it canonical- JSON serializable.
            diagnostic_actual = (
                actual.isoformat() if hasattr(actual, "isoformat") else actual
            )
            failures.append(
                {
                    "request_id": entry["request_id"],
                    "field": mapped["field"],
                    "sheet": sheet_name,
                    "cell": mapped["cell"],
                    "expected": mapped["value"],
                    "actual": diagnostic_actual,
                }
            )
    return tied, failures


def _validate_mapped_file(run_dir: Path, entry: dict) -> tuple[int, list[dict]]:
    path = run_dir / entry["path"]
    if not path.exists():
        return 0, [{"request_id": entry["request_id"], "error": "missing file"}]
    # Some auditor-controlled binary workbooks are registered artifacts but deliberately
    # expose no client-data lineage cells. An empty line_map is not a request to decode
    # the XLSX container as UTF-8 text.
    if entry.get("line_map") and path.suffix.lower() in {
        ".json",
        ".txt",
        ".csv",
    }:
        return _validate_line_map(path, entry)
    if entry.get("cell_map"):
        return _validate_cell_map(path, entry)
    return 0, []
