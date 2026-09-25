"""Full/activity journal population and delivered-export reconciliation."""

from __future__ import annotations

from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook


def run_journal_export_population_checks(
    run_dir: Path, file_map: list[dict], world: dict
) -> list[dict[str, str]]:
    """Reconcile the delivered JE export by entry and immutable line identity."""
    entry = next(
        (row for row in file_map if row.get("request_id") == "DOC-JOURNAL-EXPORT"),
        None,
    )
    if entry is None:
        return [{"error": "complete journal-entry export is missing"}]
    path = run_dir / str(entry["path"])
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception as error:
        return [{"error": f"complete journal-entry export is unreadable: {error}"}]
    try:
        header = None
        data_rows = None
        for sheet in workbook.worksheets:
            rows = list(sheet.iter_rows(values_only=True))
            for index, row in enumerate(rows):
                labels = {str(value).strip() for value in row if value is not None}
                if {"Entry number", "Line ID", "Parent line ID"} <= labels:
                    header = [
                        str(value).strip() if value is not None else "" for value in row
                    ]
                    data_rows = rows[index + 1 :]
                    break
            if header is not None:
                break
        if header is None or data_rows is None:
            return [
                {"error": "complete journal-entry export lacks stable line-ID columns"}
            ]
        columns = {label: index for index, label in enumerate(header) if label}
        actual = []
        for row in data_rows:
            entry_number = row[columns["Entry number"]]
            line_id = row[columns["Line ID"]]
            if entry_number in (None, "") or line_id in (None, ""):
                continue
            actual.append(
                {
                    "entry_number": str(entry_number),
                    "line_id": str(line_id),
                    "parent_line_id": str(row[columns["Parent line ID"]] or ""),
                    "debit": str(row[columns["Debit"]] or "0"),
                    "credit": str(row[columns["Credit"]] or "0"),
                }
            )
    finally:
        workbook.close()

    expanded = dict(world)
    for node_id, rows in list(expanded.items()):
        if (
            isinstance(rows, list)
            and rows
            and isinstance(rows[0], dict)
            and "book_layer" in rows[0]
        ):
            expanded[node_id] = [
                row for row in rows if row.get("book_layer") == "client_book"
            ]
    public_entries = {
        str(row["journal_entry_id"]): str(row["public_journal_entry_id"])
        for row in expanded.get("journal_entry") or []
    }
    expected = {
        (
            public_entries[str(row["journal_entry_id"])],
            str(row["journal_entry_line_id"]),
        ): (
            str(
                row.get("parent_journal_entry_line_id") or row["journal_entry_line_id"]
            ),
            Decimal(str(row.get("debit_amount") or "0")),
            Decimal(str(row.get("credit_amount") or "0")),
        )
        for row in expanded.get("journal_entry_line") or []
    }
    delivered = {
        (row["entry_number"], row["line_id"]): (
            row["parent_line_id"],
            Decimal(row["debit"]),
            Decimal(row["credit"]),
        )
        for row in actual
    }
    if len(delivered) != len(actual) or delivered != expected:
        return [
            {
                "error": "delivered journal export disagrees with canonical line population",
                "detail": (
                    f"expected={len(expected)}, delivered={len(actual)}, "
                    f"identity difference={sorted(set(expected) ^ set(delivered))[:10]}"
                ),
            }
        ]
    return []
