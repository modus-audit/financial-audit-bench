"""Extract bounded submitted evidence, excluding template content and grading feedback."""

from __future__ import annotations

import json
from dataclasses import dataclass, replace
from functools import lru_cache
from typing import Any

from openpyxl.formula.tokenizer import Tokenizer
from openpyxl.utils import get_column_letter, range_boundaries

from .evidence_contract import validate_cell_count
from .lookup import (
    AmbiguousLookupError,
    LookupStatus,
    MissingColumnError,
    answer_header_column,
    resolve_answer,
    table_regions,
)
from .rubric import (
    AnswerSelector,
    Check,
    FullTableSelector,
    extract_checks,
)

# Submitted evidence, excluding supplied template content and grading feedback


@lru_cache(maxsize=2048)
def _formula_structure(value: str) -> tuple:
    # Excel rewrites references when rows move. That alone is not new work.
    try:
        return tuple(
            (t.type, t.subtype, None if t.subtype == "RANGE" else t.value)
            for t in Tokenizer(value).items
        )
    except (ValueError, TypeError):
        return (value,)


def _input_cells(sheet: Any):
    """Yield populated cells and original comments, excluding grading feedback."""
    for row in sheet:
        for cell in row:
            comment = submitted_comment_text(cell.comment)
            if (cell.value is not None and str(cell.value).strip()) or (
                comment or ""
            ).strip():
                yield cell, comment


def _input_layout(sheet: Any):
    """Inputs and formula operations in order, independent of blank-row spacing."""
    for cell, comment in _input_cells(sheet):
        value = (
            _formula_structure(cell.value)
            if cell.data_type == "f"
            else json.dumps(cell.value, default=str, ensure_ascii=False)
        )
        yield cell.column, cell.data_type, value, comment


def template_content(workbook: Any) -> dict:
    """Track supplied constants/comments independently of later row movement."""
    result = {"_workbook": workbook}
    for sheet in workbook:
        values, fixed, formulas, formula_shapes = set(), {}, {}, {}
        for row in sheet:
            for cell in row:
                if cell.data_type == "f":
                    formulas[cell.coordinate] = cell.value
                    formula_shapes.setdefault(cell.column, set()).add(
                        _formula_structure(cell.value)
                    )
                if cell.value is not None and cell.data_type != "f":
                    fixed[cell.coordinate] = json.dumps(
                        cell.value, default=str, ensure_ascii=False
                    )
                    if isinstance(cell.value, str) and len(cell.value) >= 80:
                        values.add(fixed[cell.coordinate])
        result[sheet.title] = {
            "values": values,
            "fixed": fixed,
            "formulas": formulas,
            "formula_shapes": formula_shapes,
            "layout": tuple(_input_layout(sheet)),
        }
    return result


def _template_only(workbook: Any, template: dict) -> bool:
    """Moving the supplied workbook does not establish any completed work."""
    for sheet in workbook:
        if is_grading_summary(sheet):
            continue
        expected = template.get(sheet.title, {}).get("layout")
        if expected is None:
            return False
        actual = _input_layout(sheet)
        if (
            any(next(actual, None) != value for value in expected)
            or next(actual, None) is not None
        ):
            return False
    return bool(workbook.sheetnames)


def has_submitted_input(
    cell: Any, workbook: Any, template: dict, seen: set | None = None
) -> bool:
    """A supplied formula is evidence only when its input chain contains an answer."""
    seen = set() if seen is None else seen
    key = (cell.parent.title, cell.coordinate)
    if key in seen:
        return False
    seen.add(key)
    supplied = template.get(cell.parent.title, {})
    if cell.data_type != "f":
        return (
            cell.value is not None
            and bool(str(cell.value).strip())
            and json.dumps(cell.value, default=str, ensure_ascii=False)
            != supplied.get("fixed", {}).get(cell.coordinate)
        )
    try:
        references = [
            t.value for t in Tokenizer(cell.value).items if t.subtype == "RANGE"
        ]
        if not references:
            # A moved supplied constant formula is still scaffolding. A newly
            # entered constant formula (including =0) is a real entered value.
            return cell.value != supplied.get("formulas", {}).get(
                cell.coordinate
            ) and _formula_structure(cell.value) not in supplied.get(
                "formula_shapes", {}
            ).get(cell.column, set())
        # Referencing empty cells cannot establish a record, even when the
        # formula was translated by a row move or rewritten by the submitter.
        for reference in references:
            title, coord = (
                reference.rsplit("!", 1)
                if "!" in reference
                else (cell.parent.title, reference)
            )
            title = title.strip("'").replace("''", "'")
            if title not in workbook.sheetnames or "[" in reference:
                continue
            c1, r1, c2, r2 = range_boundaries(coord.replace("$", ""))
            sheet = workbook[title]
            c1, r1, c2, r2 = (
                c1 or 1,
                r1 or 1,
                c2 or sheet.max_column,
                r2 or sheet.max_row,
            )
            if (c2 - c1 + 1) * (r2 - r1 + 1) > 200_000:
                continue
            if any(
                has_submitted_input(sheet.cell(r, c), workbook, template, seen)
                for r in range(r1, r2 + 1)
                for c in range(c1, c2 + 1)
            ):
                return True
    except (ValueError, TypeError):
        pass
    return False


def extract_full_table(
    check, workbook, formulas, template, *, unchanged=None, criterion=None
):
    """Keep every populated cell below the matching header, including table notes.

    Repeated header blocks are retained. Target keys guide the judge; they never
    filter records. No total/footer wording or fixed row count bounds the data.
    """
    title = check.sheet
    spec = check.selector.to_dict()
    result = {
        "id": check.id,
        "expected": criterion if criterion is not None else check.expected,
        "actual": [],
        "extraction_errors": [],
        "has_answer": False,
        "matched_rows": 0,
    }
    if title not in workbook.sheetnames or title not in formulas.sheetnames:
        result["extraction_errors"].append(f"missing sheet: {title}")
        return result
    sheet = workbook[title]
    try:
        regions = table_regions(sheet, {}, spec["required_headers"])
    except MissingColumnError as error:
        result["extraction_errors"].append(str(error))
        return result
    supplied = template.get(title, {})
    structure = []
    for index, (header, _, end) in enumerate(regions):
        if index == len(regions) - 1:
            # Recalculation may shrink the used range by dropping comment-only
            # rows. Keep their evidence from the original submitted workbook.
            end = max(end, formulas[title].max_row + 1)
        # Unheaded side columns can contain qualifications or comment-only
        # evidence. Header width must not truncate what the judge can inspect.
        first, last = 1, max(sheet.max_column, formulas[title].max_column)
        structure.append(
            {
                "header_row": header,
                "first_row": header + 1,
                "last_row": end - 1,
                "headers": {
                    get_column_letter(col): sheet.cell(header, col).value
                    for col in range(first, last + 1)
                },
            }
        )
        for row in sheet.iter_rows(
            min_row=header + 1, max_row=end - 1, min_col=first, max_col=last
        ):
            populated = False
            for cell in row:
                original = formulas[title][cell.coordinate]
                note = submitted_comment_text(original.comment)
                if cell.value is None and not note and original.data_type != "f":
                    continue
                populated = True
                item = {
                    "id": f"{title}!{cell.coordinate}",
                    "sheet": title,
                    "row": cell.row,
                    "column": str(
                        sheet.cell(header, cell.column).value or cell.column_letter
                    ),
                    "cell": cell.coordinate,
                    "value": cell.value,
                    "number_format": cell.number_format,
                }
                if note:
                    item["comment"] = note
                if original.data_type == "f":
                    item["formula"] = original.value
                template_text = isinstance(cell.value, str) and json.dumps(
                    cell.value, ensure_ascii=False
                ) in supplied.get("values", set())
                if (
                    unchanged
                    or template_text
                    or (
                        original.data_type == "f"
                        and not has_submitted_input(original, formulas, template)
                    )
                ):
                    item["context_only"] = True
                result["actual"].append(item)
            result["matched_rows"] += int(populated)
    result["full_table"] = {"sheet": title, "regions": structure}
    result["lookup"] = {"key": spec["key"], "columns": spec["columns"]}
    result["has_answer"] = any(
        not c.get("context_only") and (c["value"] is not None or c.get("comment"))
        for c in result["actual"]
    )
    return result


def extract_check(
    check: Check,
    workbook: Any,
    formulas: Any,
    template: dict,
    *,
    unchanged: bool | None = None,
    criterion: str | None = None,
) -> dict:
    """Read one answer or candidates from the same typed resolver as numeric checks."""
    if isinstance(check.selector, FullTableSelector):
        return extract_full_table(
            check,
            workbook,
            formulas,
            template,
            unchanged=unchanged,
            criterion=criterion,
        )
    parts = extract_checks(check)
    if len(parts) > 1:
        return _extract_multiple(
            check,
            parts,
            workbook,
            formulas,
            template,
            unchanged=unchanged,
            criterion=criterion,
        )
    title = check.sheet
    result = {
        "id": check.id,
        "expected": criterion if criterion is not None else check.expected,
        "actual": [],
        "extraction_errors": [],
        "has_answer": False,
        "matched_rows": 0,
    }
    if title not in workbook.sheetnames or title not in formulas.sheetnames:
        result["extraction_errors"].append(f"missing sheet: {title}")
        return result
    selector = check.selector
    sheet = workbook[title]
    answer_columns = (
        check.definition["extract"][0].get("columns", ()) if check.kind == "llm" else ()
    )
    lookup = resolve_answer(selector, sheet, answer_columns=tuple(answer_columns))
    if lookup.needs_review:
        result.update(lookup_ambiguous=True, extraction_errors=[lookup.message])
        return result
    if lookup.status == LookupStatus.MISSING:
        result["extraction_errors"].append(lookup.message)
        return result
    count = len(lookup.rows) * max(1, len(answer_columns))
    try:
        validate_cell_count(count)
    except ValueError as error:
        result.update(
            lookup_ambiguous=True,
            extraction_errors=[str(error)],
            matched_rows=len(lookup.rows),
            selected_cell_count=count,
        )
        return result
    if lookup.status == LookupStatus.CANDIDATES:
        # Metadata locates candidates; only actual contains workbook values.
        result["table"] = {
            "sheet": title,
            "header_row": lookup.header,
            "candidate_rows": list(lookup.rows),
        }
        cells = [sheet.cell(row, lookup.column) for row in lookup.rows]
        spec = selector.to_dict()
        result["lookup"] = {
            name: spec[name] for name in ("key", "key_search_columns") if name in spec
        }
    else:
        cells = [lookup.cell]
    if len(answer_columns) > 1:
        # Identity is resolved once. All selected fields come from those rows,
        # with the same template/input checks used by the single-cell path.
        try:
            actual = []
            for row in lookup.rows:
                for name in answer_columns:
                    column = answer_header_column(sheet, lookup.header, name, row)
                    cell = sheet.cell(row, column)
                    if any(
                        cell.coordinate in merged
                        and (
                            merged.min_row != row
                            or merged.max_row != row
                            or merged.min_col != column
                        )
                        for merged in sheet.merged_cells.ranges
                    ):
                        raise AmbiguousLookupError(
                            f"selected answer {cell.coordinate} is merged across record rows or named fields"
                        )
                    actual.append(
                        answer_cell(
                            replace(selector, column=name),
                            cell,
                            str(sheet.cell(lookup.header, column).value),
                            formulas,
                            template,
                            unchanged,
                        )
                    )
        except MissingColumnError as error:
            result.update(
                lookup_ambiguous=isinstance(error, AmbiguousLookupError),
                extraction_errors=[str(error)],
            )
            return result
        result["actual"] = actual
        result["matched_rows"] = len(lookup.rows)
        if lookup.status == LookupStatus.RESOLVED:
            # A unique record gets only its selected fields, never a whole table.
            result["record"] = {"sheet": title, "row": lookup.rows[0]}
            spec = selector.to_dict()
            result["lookup"] = {
                name: spec[name]
                for name in ("key", "key_search_columns")
                if name in spec
            }
    else:
        result["actual"] = [
            answer_cell(selector, cell, lookup.label, formulas, template, unchanged)
            for cell in cells
        ]
        result["matched_rows"] = len(cells)
    result["has_answer"] = any(
        not item.get("context_only")
        and item["value"] is not None
        and str(item["value"]).strip()
        for item in result["actual"]
    )
    return result


def _extract_multiple(
    check: Check,
    parts: tuple[Check, ...],
    workbook: Any,
    formulas: Any,
    template: dict,
    *,
    unchanged: bool | None,
    criterion: str | None,
) -> dict:
    """Read explicit extracts together, within one criterion's evidence budget."""
    extracts = [
        extract_check(
            part, workbook, formulas, template, unchanged=unchanged, criterion=criterion
        )
        for part in parts
    ]
    return _combine_extracts(check, parts, extracts, workbook, criterion)


def _combine_extracts(
    check: Check,
    parts: tuple[Check, ...],
    extracts: list[dict],
    workbook: Any,
    criterion: str | None,
) -> dict:
    """Attach each selected cell to its explicit row identity without expanding context."""
    result = {
        "id": check.id,
        "expected": criterion if criterion is not None else check.expected,
        "actual": [],
        "extraction_errors": [],
        "has_answer": False,
        "matched_rows": 0,
        "records": [],
    }
    for index, (part, extracted) in enumerate(zip(parts, extracts), 1):
        result["actual"].extend(extracted["actual"])
        problems = list(extracted["extraction_errors"])
        if extracted.get("lookup_ambiguous"):
            result["lookup_ambiguous"] = True
        if not problems:
            # Do not let a vertical answer merge borrow another record's result.
            for cell in extracted["actual"]:
                if any(
                    cell["cell"] in merged and merged.min_row != merged.max_row
                    for merged in workbook[cell["sheet"]].merged_cells.ranges
                ):
                    problems.append("selected answer is merged across record rows")
                    result["lookup_ambiguous"] = True
                    break
        result["extraction_errors"].extend(
            f"extract {index} ({part.sheet}): {problem}" for problem in problems
        )
        if not problems:
            cells = extracted["actual"]
            selector = part.selector.to_dict()
            for row in dict.fromkeys(cell["row"] for cell in cells):
                result["records"].append(
                    {
                        "sheet": part.sheet,
                        "row": row,
                        "lookup": {
                            name: selector[name]
                            for name in ("key", "key_search_columns", "match")
                            if name in selector
                        },
                        "candidate_cell_ids": [
                            cell["id"] for cell in cells if cell["row"] == row
                        ],
                    }
                )
    ids = [cell["id"] for cell in result["actual"]]
    if len(set(ids)) != len(ids):
        result["extraction_errors"].append(
            "explicit extracts select overlapping answer cells"
        )
        result["lookup_ambiguous"] = True
    result["matched_rows"] = len(
        {(cell["sheet"], cell["row"]) for cell in result["actual"]}
    )
    count = sum(
        part.get("selected_cell_count", len(part["actual"])) for part in extracts
    )
    if count:
        try:
            validate_cell_count(count)
        except ValueError as error:
            result.update(
                actual=[], records=[], lookup_ambiguous=True, selected_cell_count=count
            )
            result["extraction_errors"].append(str(error))
    result["has_answer"] = any(
        not cell.get("context_only")
        and cell["value"] is not None
        and str(cell["value"]).strip()
        for cell in result["actual"]
    )
    return result


def answer_cell(
    selector: AnswerSelector,
    cell: Any,
    label: str,
    formulas: Any,
    template: dict,
    unchanged: bool | None,
) -> dict:
    """Describe submitted evidence consistently for numeric and semantic checks."""
    title = cell.parent.title
    formula_cell = formulas[title][cell.coordinate]
    item = {
        "id": f"{title}!{cell.coordinate}",
        "sheet": title,
        "row": cell.row,
        "column": label,
        "cell": cell.coordinate,
        "value": cell.value,
        # Preserve the workbook's exact display metadata for semantic judges.
        # It is intentionally opaque: no scale or percentage normalization
        # belongs in extraction.
        "number_format": cell.number_format,
    }
    if cell.value is None or isinstance(cell.value, str) and not cell.value.strip():
        item["context_only"] = True
        return item
    unchanged = (
        _template_only(formulas, template) if unchanged is None else unchanged
    ) or json.dumps(cell.value, default=str, ensure_ascii=False) in template.get(
        title, {}
    ).get("values", set())
    template_workbook = template.get("_workbook")
    if template_workbook is not None and title in template_workbook.sheetnames:
        supplied_lookup = resolve_answer(selector, template_workbook[title])
        supplied_cell = (
            supplied_lookup.cell
            if supplied_lookup.status == LookupStatus.RESOLVED
            else None
        )
        if supplied_cell is not None and supplied_cell.data_type != "f":
            unchanged = unchanged or json.dumps(cell.value, default=str) == json.dumps(
                supplied_cell.value, default=str
            )
    if formula_cell.data_type == "f" and not has_submitted_input(
        formula_cell, formulas, template
    ):
        unchanged = True
    if unchanged:
        item["context_only"] = True
    return item


def submitted_values(book: Any) -> dict:
    """Formatting and grader feedback are not submitted answers."""
    return {
        (sheet.title, cell.coordinate): (cell.value, comment)
        for sheet in book
        if not is_grading_summary(sheet)
        for cell, comment in _input_cells(sheet)
    }


@dataclass(frozen=True)
class WorkbookEvidence:
    """The original formulas and one analysis of the trusted blank template."""

    formulas: Any
    template: dict
    unchanged: bool

    @classmethod
    def from_template(cls, formulas: Any, template: Any) -> WorkbookEvidence:
        supplied = template_content(template)
        unchanged = submitted_values(formulas) == submitted_values(
            template
        ) or _template_only(formulas, supplied)
        return cls(formulas, supplied, unchanged)


GRADING_COMMENT_START = "\u2063\u2060\u2063"


GRADING_COMMENT_END = "\u2063\u2060\u2060\u2063"


SUMMARY_MARKER = "FABGradingSummary"


def is_grading_summary(sheet: Any) -> bool:
    return sheet.sheet_properties.codeName == SUMMARY_MARKER


def submitted_comment_text(comment: Any) -> str | None:
    """Remove only our delimited feedback; retain original and later user notes."""
    if comment is None:
        return None
    text = comment.text
    while GRADING_COMMENT_START in text:
        start = text.index(GRADING_COMMENT_START)
        end = text.find(GRADING_COMMENT_END, start + len(GRADING_COMMENT_START))
        if end < 0:
            break
        text = text[:start] + text[end + len(GRADING_COMMENT_END) :]
    return text or None
