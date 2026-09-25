"""Resolve declared identities and table structure without deciding correctness."""

from __future__ import annotations

from dataclasses import dataclass
from decimal import Decimal
from enum import Enum
import re
from typing import Any

from openpyxl.utils import column_index_from_string

from .comparison import equivalent, identifier_matches, normalize_text
from .rubric import (
    AnswerSelector,
    FieldSelector,
)

# Answer lookup and diagnostics


class LookupStatus(str, Enum):
    RESOLVED = "resolved"
    MISSING = "missing"
    CANDIDATES = "candidates"
    AMBIGUOUS_STRUCTURE = "ambiguous_structure"


@dataclass(frozen=True)
class LookupResult:
    status: LookupStatus
    message: str = ""
    cell: Any = None
    label: str = ""
    rows: tuple[int, ...] = ()
    header: int | None = None
    end: int | None = None
    column: int | None = None
    regions: tuple[tuple[int, int], ...] = ()

    @property
    def unresolved(self) -> bool:
        return self.status in {
            LookupStatus.CANDIDATES,
            LookupStatus.AMBIGUOUS_STRUCTURE,
        }

    @property
    def needs_review(self) -> bool:
        # A repeated fixed-form label has no bounded table to send to the judge.
        return self.status == LookupStatus.AMBIGUOUS_STRUCTURE or (
            self.status == LookupStatus.CANDIDATES and self.header is None
        )


class MergedIdentitySheet:
    """Resolve vertical merged identities without changing the submitted workbook."""

    def __init__(self, sheet: Any):
        self.sheet = sheet

    def __getattr__(self, name: str) -> Any:
        return getattr(self.sheet, name)

    def __getitem__(self, key: Any) -> Any:
        return self.sheet[key]

    def cell(self, row: int, column: int) -> Any:
        cell = self.sheet.cell(row, column)
        if cell.value is None:
            for merged in self.sheet.merged_cells.ranges:
                if (
                    merged.min_col == merged.max_col == column
                    and merged.min_row <= row <= merged.max_row
                ):
                    return self.sheet.cell(merged.min_row, column)
        return cell


def header_text(value: Any) -> str:
    """Ignore a trailing workpaper cross-reference, not the field's meaning."""
    return re.sub(r"\s*\(<\.\d+>\)\s*$", "", normalize_text(value)).strip()


def headers(sheet: Any, row: int) -> dict[str, int]:
    result: dict[str, int] = {}
    for column in range(1, sheet.max_column + 1):
        header = header_text(sheet.cell(row, column).value)
        if header:
            result.setdefault(header, column)
    return result


def keyed_rows(sheet: Any, match: dict[str, Any]) -> list[int]:
    """Locate a key/value row or a short response after a named section.

    Column letters identify fields; row numbers never bound the search.
    Unlabelled responses use their nonempty-row order after a heading, so
    inserting or deleting blank rows does not change the selected response.
    """

    def matching(
        key: dict, start: int, end: int, modes: dict | None = None
    ) -> list[int]:
        modes = modes or {}
        rows = [
            row
            for row in range(start, end)
            if all(
                any(
                    equivalent(
                        sheet.cell(row, column_index_from_string(column)).value,
                        value,
                        comparison=modes.get(column, "auto"),
                        tolerance=Decimal(".01"),
                    )
                    for value in (values if isinstance(values, list) else [values])
                )
                for column, values in key.items()
            )
        ]
        return rows

    start, end = 1, sheet.max_row + 1
    for name in ("after", "before"):
        if name not in match:
            continue
        rows = matching(match[name], start, end)
        if len(rows) != 1:
            error = AmbiguousLookupError if rows else MissingColumnError
            raise error(
                f"missing or ambiguous {name} section {match[name]!r} on {sheet.title}"
            )
        if name == "after":
            start = rows[0] + 1
        else:
            end = rows[0]
    if "key" in match:
        return matching(match["key"], start, end, match.get("comparisons"))
    # A named heading is the key for an otherwise unlabelled response block.
    following = match["following"]
    column = column_index_from_string(following["column"])
    rows = [
        row
        for row in range(start, end)
        if sheet.cell(row, column).value is not None
        and str(sheet.cell(row, column).value).strip()
    ]
    index = following.get("index", 1) - 1
    return rows[index : index + 1]


class MissingColumnError(ValueError):
    pass


class AmbiguousLookupError(MissingColumnError):
    """A structural lookup has more than one valid header or section."""


class RubricContractError(ValueError):
    """Trusted rubric structure is incompatible with its supplied template."""

    def __init__(self, defects: list[dict]):
        self.defects = defects
        super().__init__(
            "Trusted rubric/template contract failed: "
            + "; ".join(
                f"{defect['criterion_id']}: {defect['message']}" for defect in defects
            )
        )


def _section_end(sheet: Any, header_row: int, end_header: Any) -> int | None:
    if end_header is None:
        return None
    boundary_labels = {header_text(name) for name in end_header}
    boundary = next(
        (
            row
            for row in range(header_row + 1, sheet.max_row + 1)
            if boundary_labels <= headers(sheet, row).keys()
        ),
        None,
    )
    if boundary is None:
        raise MissingColumnError(
            f"missing section boundary {end_header!r} on {sheet.title}"
        )
    return boundary


def _column(headers: dict[str, int], name: str, sheet: str) -> int:
    try:
        return headers[header_text(name)]
    except KeyError as error:
        raise MissingColumnError(f"missing column {name!r} on {sheet}") from error


def _header_columns(sheet: Any, header_row: int, name: str) -> list[int]:
    return [
        column
        for column in range(1, sheet.max_column + 1)
        if header_text(sheet.cell(header_row, column).value) == header_text(name)
    ]


def unique_header_column(sheet: Any, header_row: int, name: str) -> int:
    """Resolve an alternate column only within the selected table header."""
    columns = _header_columns(sheet, header_row, name)
    if len(columns) != 1:
        problem = "missing" if not columns else "ambiguous"
        error = MissingColumnError if not columns else AmbiguousLookupError
        raise error(
            f"{problem} column {name!r} on {sheet.title} header row {header_row}"
        )
    return columns[0]


def answer_header_column(sheet: Any, header_row: int, name: str, row: int) -> int:
    """Allow a duplicate answer header only when the keyed row is unambiguous.

    Identity columns and trusted-template headers still require unique names.
    Two populated alternatives remain ambiguous, even if their values agree.
    """
    columns = _header_columns(sheet, header_row, name)
    if len(columns) <= 1:
        return unique_header_column(sheet, header_row, name)
    populated = []
    for column in columns:
        cell = sheet.cell(row, column)
        for merged in sheet.merged_cells.ranges:
            if cell.coordinate in merged:
                cell = sheet.cell(merged.min_row, merged.min_col)
                break
        if cell.value is not None and str(cell.value).strip():
            populated.append(column)
    if len(populated) > 1:
        raise AmbiguousLookupError(
            f"ambiguous populated column {name!r} on {sheet.title} row {row}"
        )
    # All alternatives blank means a blank answer, not an arbitrary verdict.
    return populated[0] if populated else columns[0]


def _matching_rows(
    sheet: Any,
    headers: dict[str, int],
    header_row: int,
    key: dict[str, Any],
    *,
    comparisons: dict[str, str] | None = None,
    key_search_columns: dict[str, list[str]] | None = None,
    key_aliases: dict[str, list[str]] | None = None,
    key_alias_comparisons: dict[str, str] | None = None,
    end_row: int | None = None,
    tolerance: Decimal,
) -> list[int]:
    """Match the complete key in its own fields before searching alternate columns."""
    comparisons = comparisons or {}
    aliases = {
        name: {normalize_text(value) for value in values}
        for name, values in (key_aliases or {}).items()
    }

    def alias_matches(name: str, value: Any) -> bool:
        if (key_alias_comparisons or {}).get(name) == "identifier":
            return any(
                identifier_matches(value, alias) for alias in aliases.get(name, set())
            )
        return normalize_text(value) in aliases.get(name, set())

    key_columns = {
        _column(headers, name, sheet.title): (name, expected)
        for name, expected in key.items()
    }

    def matching_rows(columns: dict[int, list[int]]) -> list[int]:
        return [
            row
            for row in range(header_row + 1, end_row or sheet.max_row + 1)
            if all(
                any(
                    equivalent(
                        sheet.cell(row, candidate).value,
                        expected,
                        comparison=comparisons.get(name, "auto"),
                        tolerance=tolerance,
                    )
                    or alias_matches(name, sheet.cell(row, candidate).value)
                    for candidate in columns[column]
                )
                for column, (name, expected) in key_columns.items()
            )
        ]

    primary_rows = matching_rows({column: [column] for column in key_columns})
    # Preserve ambiguity between real records. A mention in a comment must
    # neither compete with a primary match nor break a tie between duplicates.
    if primary_rows or not key_search_columns:
        return primary_rows
    candidate_columns = {
        column: [
            column,
            *[
                unique_header_column(sheet, header_row, candidate)
                for candidate in (key_search_columns or {}).get(name, [])
            ],
        ]
        for column, (name, _) in key_columns.items()
    }
    return matching_rows(candidate_columns)


def _table_bounds(sheet: Any, section: dict | None) -> tuple[int, int]:
    """Scope a table by unique, explicit section labels, never row proximity."""
    start, end = 1, sheet.max_row + 1
    if section is None:
        return start, end
    for name in ("after", "before"):
        if name not in section:
            continue
        rows = keyed_rows(sheet, {"key": section[name]})
        if len(rows) != 1:
            error = AmbiguousLookupError if rows else MissingColumnError
            raise error(
                f"missing or ambiguous {name} section {section[name]!r} on {sheet.title}"
            )
        if name == "after":
            start = rows[0] + 1
        else:
            end = rows[0]
    if start >= end:
        raise MissingColumnError(f"invalid section boundaries on {sheet.title}")
    return start, end


def table_regions(
    sheet: Any,
    check: dict,
    names: list[str],
    *,
    answer_columns: tuple[str, ...] = (),
) -> list[tuple[int, dict, int]]:
    """Find bounded copies of a table, retaining every copy with the same columns."""
    start, end = _table_bounds(sheet, check.get("section"))
    names = list(dict.fromkeys([*names, *check.get("required_headers", [])]))
    required = {header_text(name) for name in names}
    if not required:
        raise ValueError("table lookup requires named headers")
    rows = [row for row in range(start, end) if required <= headers(sheet, row).keys()]
    candidates = []
    for row in rows:
        try:
            boundary = _section_end(sheet, row, check.get("end_header"))
        except MissingColumnError:
            continue  # A supporting table after the declared footer is outside this table.
        if boundary is None or boundary <= end:
            candidates.append((row, boundary or end))
    if not candidates and rows:
        raise MissingColumnError(
            f"missing section boundary {check.get('end_header')!r} inside table scope on {sheet.title}"
        )
    if not candidates:
        raise MissingColumnError(f"missing table headers {names!r} on {sheet.title}")
    layouts = [
        tuple(
            tuple(_header_columns(sheet, row, name))
            if name in answer_columns
            else (unique_header_column(sheet, row, name),)
            for name in names
        )
        for row, _ in candidates
    ]
    if len(set(layouts)) != 1:
        raise AmbiguousLookupError(
            f"inconsistent table columns {names!r} on {sheet.title}"
        )
    return [
        (
            row,
            headers(sheet, row),
            min(boundary, candidates[i + 1][0])
            if i + 1 < len(candidates)
            else boundary,
        )
        for i, (row, boundary) in enumerate(candidates)
    ]


def table_header(sheet: Any, check: dict, names: list[str]) -> tuple[int, dict, int]:
    """Require one table when validating the trusted template's contract."""
    regions = table_regions(sheet, check, names)
    if len(regions) != 1:
        raise AmbiguousLookupError(
            f"ambiguous table headers {names!r} on {sheet.title}"
        )
    return regions[0]


def table_record_rows(
    sheet: Any, check: dict, names: list[str]
) -> tuple[int, tuple[int, ...]]:
    regions = table_regions(sheet, check, names)
    return regions[0][0], tuple(
        row for header, _, end in regions for row in range(header + 1, end)
    )


def _answer_headers(selector: dict) -> list[str]:
    return [
        *selector["key"],
        selector["column"],
        *(
            name
            for columns in selector.get("key_search_columns", {}).values()
            for name in columns
        ),
    ]


def _is_merged_table_note(
    sheet: Any,
    regions: list[tuple[int, dict, int]],
    row: int,
) -> bool:
    """A note spanning every table field is not a tabular record.

    Do not infer this from label wording: a differently worded record can carry
    a conflicting answer and must remain a candidate.
    """
    header = next(h for h, _, end in regions if h < row < end)
    columns = [cell.column for cell in sheet[header] if normalize_text(cell.value)]
    return any(
        merged.min_row == merged.max_row == row
        and merged.min_col < merged.max_col
        and all(merged.min_col <= col <= merged.max_col for col in columns)
        for merged in sheet.merged_cells.ranges
    )


def resolve_answer(
    selector: AnswerSelector,
    sheet: Any,
    tolerance: Decimal = Decimal("0.01"),
    *,
    answer_columns: tuple[str, ...] = (),
) -> LookupResult:
    """Resolve identity once; return cardinality separately from evaluation policy."""
    check = selector.to_dict()
    header, end = None, None
    regions = []
    try:
        if isinstance(selector, FieldSelector):
            rows = keyed_rows(MergedIdentitySheet(sheet), check["match"])
            column = column_index_from_string(selector.column_letter)
            label = selector.column
        else:
            key = check["key"]
            # Only answer headers may be duplicated; identity/search fields
            # retain their strict uniqueness requirements.
            identity_headers = [
                *key,
                *(n for ns in check.get("key_search_columns", {}).values() for n in ns),
            ]
            regions = table_regions(
                sheet,
                check,
                _answer_headers(check),
                answer_columns=tuple(
                    name
                    for name in (answer_columns or (selector.column,))
                    if name not in identity_headers
                ),
            )
            identities = MergedIdentitySheet(sheet)
            rows = [
                row
                for region_header, columns, region_end in regions
                for row in _matching_rows(
                    identities,
                    columns,
                    region_header,
                    key,
                    comparisons=check.get("key_comparisons", {}),
                    key_search_columns=check.get("key_search_columns"),
                    key_aliases=check.get("key_aliases"),
                    key_alias_comparisons=check.get("key_alias_comparisons"),
                    end_row=region_end,
                    tolerance=tolerance,
                )
            ]
            rows = [
                row for row in rows if not _is_merged_table_note(sheet, regions, row)
            ]
            header, _, _ = regions[0]
            end = regions[-1][2]
            answer_columns = {
                answer_header_column(sheet, header, selector.column, row)
                for row in rows
            }
            if len(answer_columns) > 1:
                raise AmbiguousLookupError(
                    f"inconsistent answer columns {selector.column!r} on {sheet.title}"
                )
            column = (
                next(iter(answer_columns))
                if answer_columns
                else _column(headers(sheet, header), selector.column, sheet.title)
            )
            label = str(sheet.cell(header, column).value)
    except AmbiguousLookupError as error:
        return LookupResult(LookupStatus.AMBIGUOUS_STRUCTURE, str(error))
    except MissingColumnError as error:
        return LookupResult(LookupStatus.MISSING, str(error))
    if len(rows) > 1:
        return LookupResult(
            LookupStatus.CANDIDATES,
            f"ambiguous answer row on {sheet.title}: matched {len(rows)} rows",
            label=label,
            rows=tuple(rows),
            header=header,
            end=end,
            column=column,
            regions=tuple((h, e) for h, _, e in regions),
        )
    if not rows:
        return LookupResult(
            LookupStatus.MISSING, f"missing answer row on {sheet.title}: matched 0 rows"
        )
    cell = sheet.cell(rows[0], column)
    for merged in sheet.merged_cells.ranges:
        if cell.coordinate in merged:
            cell = sheet.cell(merged.min_row, merged.min_col)
            break
    return LookupResult(
        LookupStatus.RESOLVED,
        cell=cell,
        label=label,
        rows=tuple(rows),
        header=header,
        end=end,
        column=column,
    )
