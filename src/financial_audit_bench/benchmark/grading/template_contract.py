"""Validate rubric selectors against trusted, unmodified workbook templates."""

from __future__ import annotations

from typing import Any

from openpyxl.utils import column_index_from_string

from .lookup import (
    MissingColumnError,
    RubricContractError,
    _answer_headers,
    keyed_rows,
    table_header,
    unique_header_column,
)
from .rubric import (
    Check,
    extract_checks,
    lookup_details,
)

# Trusted template validation and navigation


def scan_columns(check: dict, names: list[str]) -> list[str]:
    summary = check.get("summary_rows", {})
    return list(
        dict.fromkeys(
            names
            + ([summary["column"], *summary["identity_columns"]] if summary else [])
        )
    )


def contract_anchor(check: Check, sheet: Any) -> dict | None:
    """Resolve declared structure, never requiring an expected record to exist.

    The returned location is navigation context, not an answer or grading target.
    The same contract is checked on the trusted template before evaluation and
    on the submission when a missing answer needs a safe section/header link.
    """
    selector = check.selector.to_dict() if check.selector else check.to_dict()
    if "match" in selector:
        # This validates bounded labels and selector syntax while allowing an
        # empty result: agents are permitted to add records and response rows.
        match = selector["match"]
        keyed_rows(sheet, match)
        for boundary in ("after", "before"):
            if boundary not in match:
                continue
            rows = keyed_rows(sheet, {"key": match[boundary]})
            if len(rows) == 1:
                column = column_index_from_string(next(iter(match[boundary])))
                return {
                    "sheet": sheet.title,
                    "cell": sheet.cell(rows[0], column).coordinate,
                    "role": "section_header",
                }
        return None
    if selector.get("kind") == "table":
        names = selector["required_headers"]
    elif selector.get("kind") == "selection_coverage":
        names = scan_columns(
            selector, selector["key_columns"] + selector["completion_columns"]
        )
    elif selector.get("kind") == "notes_have_context":
        names = scan_columns(
            selector, selector["primary_columns"] + selector["note_columns"]
        )
    elif selector.get("kind") == "row_count":
        names = scan_columns(selector, selector["key_columns"])
    elif selector.get("kind") == "date_window":
        names = scan_columns(
            selector, [selector["date_column"], *selector["record_columns"]]
        )
    else:
        names = _answer_headers(selector)
    header, _, _ = table_header(sheet, selector, names)
    column = unique_header_column(sheet, header, selector.get("column", names[0]))
    return {
        "sheet": sheet.title,
        "cell": sheet.cell(header, column).coordinate,
        "role": "table_header",
    }


def validate_template_contract(checks: tuple[Check, ...], template: Any) -> None:
    """Fail consistently when trusted selectors cannot locate template structure."""
    defects = []
    for check in (part for declared in checks for part in extract_checks(declared)):
        try:
            if check.sheet not in template.sheetnames:
                raise MissingColumnError(
                    f"missing declared worksheet {check.sheet!r} in trusted template"
                )
            contract_anchor(check, template[check.sheet])
        except (MissingColumnError, ValueError, KeyError) as error:
            defects.append(
                {
                    "defect_id": "template_contract:" + check.id,
                    "criterion_id": check.id,
                    "sheet": check.sheet,
                    "selector": lookup_details(check.to_dict()),
                    "message": str(error),
                }
            )
    if defects:
        raise RubricContractError(defects)


def missing_answer_context(check: Check, sheet: Any) -> list[dict]:
    try:
        anchor = contract_anchor(check, sheet)
    except (MissingColumnError, ValueError, KeyError):
        return []
    return [anchor] if anchor else []
