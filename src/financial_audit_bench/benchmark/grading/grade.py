"""Run the shared grader and write Harbor reports, rewards, and feedback."""

from __future__ import annotations

import argparse
import hashlib
import json
import math
import stat
from collections import defaultdict
from contextlib import ExitStack
from decimal import Decimal
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from . import deterministic, feedback, judge, template_contract
from .lookup import (
    RubricContractError,
)
from .recalculate import (
    recalculated_workbooks,
)
from .recalculation_cache import RecalculationCache
from .results import (
    check_result,
    score_summary,
    write_json,
)
from .rubric import (
    Check,
    compile_checks,
    load_rubric,
    lookup_details,
)
from .workbook import (
    WorkbookEvidence,
)


def _missing_output_checks(
    rubric: dict[str, Any], message: str | None = None
) -> list[dict[str, Any]]:
    return [
        check_result(
            check, tab["sheet"], False, message or f"missing output: {rubric['output']}"
        )
        for tab in rubric["tabs"]
        for check in tab["checks"]
    ]


def _contract_error_checks(rubric: dict, error: RubricContractError) -> list[dict]:
    defects = {item["criterion_id"]: item for item in error.defects}
    results = []
    for tab in rubric["tabs"]:
        for check in tab["checks"]:
            defect = defects.get(check["id"])
            results.append(
                check_result(
                    check,
                    tab["sheet"],
                    None,
                    defect["message"]
                    if defect
                    else "Evaluation not run because the trusted rubric/template contract is defective.",
                    status="error",
                    error="rubric_contract_error",
                    defect_ids=[defect["defect_id"]]
                    if defect
                    else [d["defect_id"] for d in error.defects],
                )
            )
    return results


def _tab_criteria(
    rubric: dict[str, Any], checks: list[dict[str, Any]]
) -> list[dict[str, Any]]:
    by_sheet: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for check in checks:
        by_sheet[check["sheet"]].append(check)
    criteria = []
    for tab in rubric["tabs"]:
        rows = by_sheet[tab["sheet"]]
        summary = score_summary(rows)
        criteria.append(
            {
                "id": tab["id"],
                "sheet": tab["sheet"],
                **summary,
                "message": "Evaluation error; no score"
                if summary["score"] is None
                else f"{summary['passed_checks']} of {summary['total_checks']} checks passed",
            }
        )
    return criteria


def _hash(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _rubric_hashes(task: Path) -> dict:
    return {"rubric.json": _hash(task / "rubric.json")}


def regenerate_feedback(workspace: Path, task: Path, output: Path) -> dict:
    """Retry presentation from a saved report; never recalculate or call a judge."""
    report = json.loads(output.read_text())
    rubric = load_rubric(task, automatic_only=report.get("score_scope") == "automatic")
    source = workspace / rubric["output"]
    provenance = report.get("provenance", {})
    if (
        source.is_symlink()
        or not source.is_file()
        or provenance.get("submission_sha256") != _hash(source)
        or provenance.get("rubrics_sha256") != _rubric_hashes(task)
    ):
        raise ValueError(
            "Feedback requires the same submission and rubric files as the saved report"
        )
    ids = {check["id"] for tab in rubric["tabs"] for check in tab["checks"]}
    if {check["id"] for check in report.get("checks", [])} != ids or len(
        report["checks"]
    ) != len(ids):
        raise ValueError("Saved report does not cover the rubric")
    feedback.export_feedback(report, source, rubric, output)
    return report


def _evaluate_checks(
    compiled: list[Check],
    actual: Any,
    evidence: WorkbookEvidence,
    output: Path,
    *,
    tolerance: Decimal,
    automatic_only: bool,
) -> tuple[list[dict], dict | None]:
    """Evaluate each check with its declared method; never substitute a judge."""
    judge_summary = None
    semantic_rules = [check for check in compiled if check.kind == "llm"]
    checks = []
    for check in compiled:
        if check.kind == "llm":
            continue
        if check.sheet in actual.sheetnames:
            evaluation = deterministic.evaluate_check(
                check,
                actual[check.sheet],
                tolerance=tolerance,
                evidence=evidence,
            )
            checks.append(evaluation.report)
        else:
            checks.append(
                check_result(
                    check.to_dict(),
                    check.sheet,
                    False,
                    f"missing sheet: {check.sheet}",
                )
            )
    if semantic_rules and not automatic_only:
        judged, judge_summary = judge.evaluate_checks(
            semantic_rules,
            actual,
            evidence,
            output,
        )
        checks.extend(judged)
    return checks, judge_summary


def evaluate(
    workspace: Path,
    task: Path,
    template_path: Path,
    output: Path,
    *,
    automatic_only: bool = False,
    recalculation_cache: Path | None = None,
) -> dict[str, Any]:
    """Grade with one validated contract and one owned set of workbook objects."""
    rubric = load_rubric(task, automatic_only=automatic_only)
    rules = {c["id"]: c for tab in rubric["tabs"] for c in tab["checks"]}
    compiled = compile_checks(rubric)
    compiled_by_id = {check.id: check for check in compiled}
    filename = rubric["output"]
    actual_path = workspace / filename
    if actual_path.is_symlink() or (
        actual_path.exists() and not stat.S_ISREG(actual_path.stat().st_mode)
    ):
        raise ValueError("submission must be a regular file")
    output_exists = actual_path.is_file()
    tolerance = Decimal(str(rubric["comparison"]["numeric_absolute_tolerance"]))
    judge_summary, preflight_error = None, None
    with ExitStack() as stack:

        def open_book(path: Path, *, data_only: bool = False) -> Any:
            book = load_workbook(path, data_only=data_only)
            stack.callback(book.close)
            return book

        template = open_book(template_path)
        try:
            template_contract.validate_template_contract(compiled, template)
        except RubricContractError as error:
            preflight_error = error
        if preflight_error is not None:
            checks = _contract_error_checks(rubric, preflight_error)
        elif not output_exists:
            checks = _missing_output_checks(rubric)
        else:
            formulas = open_book(actual_path)
            evidence = WorkbookEvidence.from_template(formulas, template)
            if evidence.unchanged:
                checks = _missing_output_checks(
                    rubric, "template contains no submitted answers"
                )
            else:
                if recalculation_cache is not None:
                    recalculated_path = RecalculationCache(recalculation_cache).resolve(
                        actual_path
                    )
                else:
                    recalculated = stack.enter_context(
                        recalculated_workbooks({"submission": actual_path})
                    )
                    recalculated_path = recalculated["submission"]
                actual = open_book(recalculated_path, data_only=True)
                checks, judge_summary = _evaluate_checks(
                    compiled,
                    actual,
                    evidence,
                    output,
                    tolerance=tolerance,
                    automatic_only=automatic_only,
                )
            for result in checks:
                if result["passed"] is False and not result.get("matched_locations"):
                    rule = rules[result["id"]]
                    title = (
                        rule["extract"][0]["sheet"]
                        if rule["kind"] == "llm"
                        else result["sheet"]
                    )
                    if title in formulas.sheetnames:
                        result["context_locations"] = (
                            template_contract.missing_answer_context(
                                compiled_by_id[result["id"]], formulas[title]
                            )
                        )
    if len(checks) != len(rules) or {check["id"] for check in checks} != rules.keys():
        raise ValueError("grader did not evaluate every rubric check")
    for result in checks:
        result["selector"] = lookup_details(rules[result["id"]])
    numeric = [c for c in checks if c["kind"] != "llm"]
    judged = [c for c in checks if c["kind"] == "llm"]

    def counts(items: list[dict]) -> dict:
        return {
            key: value
            for key, value in score_summary(items).items()
            if key
            in {"passed_checks", "total_checks", "earned_points", "possible_points"}
        }

    grader_files = [
        *sorted(Path(__file__).parent.glob("*.py")),
        Path(__file__).parents[2] / "llms.py",
    ]
    config = {
        "model": judge.JUDGE_MODEL,
        "reasoning_effort": judge.JUDGE_REASONING_EFFORT,
        "workers": judge.JUDGE_WORKERS,
        "request_timeout_seconds": judge.JUDGE_REQUEST_TIMEOUT_SECONDS,
    }
    provenance = {
        "rubrics_sha256": _rubric_hashes(task),
        "template_sha256": _hash(template_path),
        "submission_sha256": _hash(actual_path) if output_exists else None,
        "judge_config": config,
        "judge_config_sha256": hashlib.sha256(
            json.dumps(config, sort_keys=True).encode()
        ).hexdigest(),
        "grader_files_sha256": {path.name: _hash(path) for path in grader_files},
    }
    report = {
        **score_summary(checks),
        "checks": checks,
        "criteria": _tab_criteria(rubric, checks),
        "score_scope": "automatic" if automatic_only else "automatic_and_semantic",
        "automatic": counts(numeric),
        "semantic": {**counts(judged), "judge": judge_summary},
        "provenance": provenance,
        # Retained for existing report consumers; provenance names all inputs explicitly.
        "rubric_sha256": provenance["rubrics_sha256"]["rubric.json"],
    }
    if report["score"] is None:
        report["evaluation_outcome"] = "evaluation_error"
    if judge_summary and judge_summary["status"] == "error":
        report.update(score=None, passed=None, evaluation_outcome="judge_error")
    if preflight_error is not None:
        report.update(
            evaluation_outcome="rubric_contract_error",
            error={
                "type": "RubricContractError",
                "message": str(preflight_error),
                "defects": preflight_error.defects,
            },
        )
    write_json(output, report)
    if output_exists:
        feedback.export_feedback(report, actual_path, rubric, output)
    return report


def main() -> None:
    parser = argparse.ArgumentParser(description="Grade an audit workpaper.")
    parser.add_argument("--workspace", type=Path, required=True)
    parser.add_argument(
        "--task", type=Path, required=True, help="Trusted rubric directory."
    )
    parser.add_argument("--template", type=Path)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument(
        "--annotate-only",
        action="store_true",
        help="Regenerate feedback from an existing report; no grading calls.",
    )
    parser.add_argument("--automatic-only", action="store_true")
    parser.add_argument(
        "--recalculation-cache",
        type=Path,
        help="Host-prepared batch cache; missing or invalid entries are evaluation errors.",
    )
    parser.add_argument(
        "--reward", action="store_true", help="Also write Harbor's reward.json."
    )
    args = parser.parse_args()
    if not args.annotate_only and args.template is None:
        parser.error("grading requires --template")
    args.output.parent.mkdir(parents=True, exist_ok=True)
    reward_path = args.output.parent / "reward.json"
    if args.reward and not args.annotate_only:
        reward_path.unlink(missing_ok=True)
        # Harbor falls back to reward.txt when reward.json is absent.
        reward_path.with_suffix(".txt").unlink(missing_ok=True)
    try:
        if args.annotate_only:
            report = regenerate_feedback(args.workspace, args.task, args.output)
        else:
            report = evaluate(
                args.workspace,
                args.task,
                args.template,
                args.output,
                automatic_only=args.automatic_only,
                recalculation_cache=args.recalculation_cache,
            )
    except Exception as error:
        if args.annotate_only:
            raise SystemExit(f"Feedback regeneration failed: {error}") from error
        report = {
            "score": None,
            "passed": None,
            "evaluation_outcome": "error",
            "error": {"type": type(error).__name__, "message": str(error)},
        }
    write_json(args.output, report)
    if args.annotate_only:
        if report["workbook_highlighting"]["status"] != "complete":
            raise SystemExit("Feedback regeneration failed; see " + str(args.output))
        if not args.reward:
            return
    score = report.get("score")
    if (
        not isinstance(score, (int, float))
        or not math.isfinite(score)
        or not 0 <= score <= 1
    ):
        raise SystemExit("Evaluation failed; see " + str(args.output))
    if args.reward:
        write_json(reward_path, {"reward": score})


if __name__ == "__main__":
    main()
