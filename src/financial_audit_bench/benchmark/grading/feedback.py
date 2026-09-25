"""Render and export annotated workbook copies without changing grading evidence."""

from __future__ import annotations

import shutil
import tempfile
from collections import defaultdict
from copy import copy
from decimal import Decimal
from pathlib import Path
from textwrap import wrap
from typing import Any

from openpyxl import load_workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import quote_sheetname
from openpyxl.worksheet.hyperlink import Hyperlink
from openpyxl.worksheet.page import PageMargins
from openpyxl.worksheet.views import Selection

from .results import (
    score_summary,
    write_json,
)
from .rubric import (
    compile_check,
    display_value,
    expectation,
    lookup_details,
)
from .template_contract import (
    missing_answer_context,
)
from .workbook import (
    GRADING_COMMENT_END,
    GRADING_COMMENT_START,
    SUMMARY_MARKER,
    is_grading_summary,
    submitted_comment_text,
)

# Invisible Unicode boundaries keep feedback out of submitted answer evidence
# without displaying banners.

CHECK_COLORS = {
    "deterministic": {"passed": "66E0FF", "failed": "FF78BD"},
    "llm": {"passed": "C5F05B", "failed": "B58AFF"},
}


COLOR_LEGEND = (
    "Colors: deterministic pass = cyan, fail = pink; LLM pass = lime, fail = violet."
)


SHARED_CELL_RULE = "Shared cells show failures first; deterministic checks take priority when results match."


GRADING_FORMULA = 'N("Financial Audit Bench grading")=0'
UNHIGHLIGHTED_PASSES = {"notes_have_context", "row_count", "date_window"}
TABLE_CHECKS = UNHIGHLIGHTED_PASSES | {"selection_coverage"}


def _summary_link(
    cell: Any, title: str, coordinate: str, *, bold: bool = False
) -> None:
    cell.hyperlink = Hyperlink(
        ref=cell.coordinate,
        location=f"{quote_sheetname(title)}!{coordinate}",
    )
    cell.font = Font(
        name="Helvetica Neue",
        size=11,
        color="245A81",
        underline="single",
        bold=bold,
    )


def _unlocated_label(rule: dict, result: dict) -> str:
    if result.get("status") == "needs_review":
        return "Unresolved lookup — needs review"
    if rule.get("kind") in TABLE_CHECKS and type(result.get("passed")) is bool:
        return "Table check passed" if result["passed"] else "Table check failed"
    return (
        "Answer not located — scored 0"
        if result.get("passed") is False
        else "Evaluation error — no score"
    )


def _verdict_label(rule: dict, result: dict) -> str:
    if rule.get("kind") == "notes_have_context" and result.get("passed") is False:
        return "Wrong (shared)"
    if result.get("status") == "needs_review":
        return "Needs review"
    if result.get("passed") is True:
        return "Correct"
    if result.get("passed") is False:
        return "Wrong (0)"
    return "No score"


def add_grading_summary(
    workbook: Any, rubric: dict, checks: list[dict], comments: dict
) -> dict:
    """Add a score and a complete check index with internal links to graded cells."""
    for sheet in list(workbook):
        if is_grading_summary(sheet):
            workbook.remove(sheet)
    sheet = workbook.create_sheet("Grading Summary", 0)
    sheet.sheet_properties.codeName = SUMMARY_MARKER
    sheet.sheet_properties.tabColor = "233B53"
    sheet.sheet_view.showGridLines = False
    sheet.sheet_view.zoomScale = 90
    for column, width in {"A": 5.5, "B": 62, "C": 31, "D": 10, "E": 14}.items():
        sheet.column_dimensions[column].width = width
    sheet.sheet_format.defaultRowHeight = 20
    rules = {c["id"]: c for tab in rubric["tabs"] for c in tab["checks"]}
    results = [c for c in checks if c["id"] in rules]
    totals = score_summary(results)
    passed, score = totals["passed_checks"], totals["score"] if results else None
    locations: dict[str, list[tuple[str, str]]] = defaultdict(list)
    for location, explanations in comments.items():
        for check_id in explanations:
            locations[check_id].append(location)

    font = "Helvetica Neue"
    ink, muted, line = "233B53", "607286", "DDE5ED"

    def text(
        row: int,
        value: Any,
        start: int = 1,
        end: int = 5,
        *,
        color: str = "FFFFFF",
        size: float = 11,
        bold: bool = False,
        height: float = 26,
        align: str = "left",
        font_color: str = ink,
    ) -> Any:
        if start != end:
            sheet.merge_cells(
                start_row=row, start_column=start, end_row=row, end_column=end
            )
        cell = sheet.cell(row, start, value)
        cell.font = Font(name=font, size=size, bold=bold, color=font_color)
        cell.alignment = Alignment(
            horizontal=align,
            vertical="center",
            wrap_text=True,
            indent=1 if align in {"left", "right"} else 0,
        )
        for column in range(start, end + 1):
            sheet.cell(row, column).fill = PatternFill("solid", fgColor="FF" + color)
        sheet.row_dimensions[row].height = height
        return cell

    def tint(color: str) -> str:
        return "".join(
            f"{round(int(color[i : i + 2], 16) * 0.22 + 255 * 0.78):02X}"
            for i in (0, 2, 4)
        )

    def row_height(value: str, width: int) -> float:
        lines = sum(max(1, len(wrap(part, width=width))) for part in value.split("\n"))
        return max(30, 14 * lines + 12)

    text(1, "Grading summary", size=22, bold=True, height=36)
    text(2, rubric["output"], size=10, font_color=muted, height=24)
    for cell in sheet[2]:
        cell.border = Border(bottom=Side(style="thin", color=line))
    sheet.row_dimensions[3].height = 10
    text(4, "Overall score", end=2, bold=True, color="F0F4F8", height=40)
    text(
        4,
        score if score is not None else "Unavailable",
        3,
        5,
        size=26,
        bold=True,
        color="F0F4F8",
        align="right",
        height=40,
    ).number_format = "0.00%"
    text(
        5,
        "Points earned / possible",
        end=2,
        color="F0F4F8",
        font_color=muted,
        height=28,
    )
    points = totals["earned_points"] if score is not None else "Unavailable"
    text(
        5,
        f"{points} / {totals['possible_points']}",
        3,
        5,
        color="F0F4F8",
        bold=True,
        align="right",
        height=28,
    )
    text(
        6,
        "Grading method",
        end=2,
        color=ink,
        font_color="FFFFFF",
        size=10,
        bold=True,
        height=24,
    )
    text(
        6,
        "Points / possible",
        3,
        3,
        color=ink,
        font_color="FFFFFF",
        size=10,
        bold=True,
        align="center",
        height=24,
    )
    text(
        6,
        "Score",
        4,
        5,
        color=ink,
        font_color="FFFFFF",
        size=10,
        bold=True,
        align="right",
        height=24,
    )
    for row, kind, label in ((7, "deterministic", "Deterministic"), (8, "llm", "LLM")):
        subset = [c for c in results if check_type(rules[c["id"]], c) == kind]
        summary = score_summary(subset)
        count = (
            summary["earned_points"] if summary["score"] is not None else "Unavailable"
        )
        color = "F6F8FA" if row == 8 else "FFFFFF"
        text(row, label, end=2, bold=True, color=color)
        text(
            row,
            f"{count} / {summary['possible_points']}",
            3,
            3,
            color=color,
            align="center",
        )
        rate = summary["score"] if subset and summary["score"] is not None else "N/A"
        text(row, rate, 4, 5, color=color, align="right").number_format = "0.00%"
    sheet.row_dimensions[9].height = 10
    text(
        10,
        "Cell links open graded answers; header links provide context only. "
        "Empty-row documentation checks share one point per workpaper.",
        size=9,
        font_color=muted,
        height=30,
    )
    text(11, SHARED_CELL_RULE, size=9, font_color=muted, height=20)
    text(12, COLOR_LEGEND, size=9, font_color=muted, height=20)
    sheet.row_dimensions[13].height = 12
    groups = [
        (
            kind,
            outcome,
            [
                c
                for c in results
                if check_type(rules[c["id"]], c) == kind and c.get("passed") is verdict
            ],
        )
        for kind in ("deterministic", "llm")
        for outcome, verdict in (("failed", False), ("passed", True))
    ]
    ungraded = [c for c in results if type(c.get("passed")) is not bool]
    if ungraded:
        groups.append(("ungraded", "ungraded", ungraded))
    row, linked_cells, context_links = 14, 0, 0
    for kind, outcome, items in groups:
        color = CHECK_COLORS[kind][outcome] if kind != "ungraded" else "E7E6E6"
        label = ("LLM" if kind == "llm" else "Deterministic") + ": " + outcome.title()
        unit = "check" if len(items) == 1 else "checks"
        text(
            row,
            f"{label if kind != 'ungraded' else 'Not graded'} ({len(items)} {unit})",
            color="EDF2F7",
            bold=True,
            size=12,
            height=32,
        ).border = Border(left=Side(style="thick", color=color))
        row += 1
        for column, label in enumerate(
            ("#", "Criterion", "Worksheet", "Cell", "Result"), 1
        ):
            text(
                row,
                label,
                column,
                column,
                color="F7F9FB",
                font_color=muted,
                size=10,
                bold=True,
                align="center" if column in (1, 4, 5) else "left",
                height=24,
            )
        row += 1
        if not items:
            text(row, "No checks in this group.", size=11, font_color=muted)
            row += 1
        for number, result in enumerate(items, 1):
            rule = rules[result["id"]]
            description = (
                rule.get("description") or result.get("description") or result["id"]
            )
            if rule.get("kind") == "notes_have_context":
                description += "\nShared one-point criterion across all tables."
            context = None
            if not locations.get(result["id"]) and not (
                rule.get("kind") in UNHIGHLIGHTED_PASSES
                and result.get("passed") is True
            ):
                details = result.get("selector") or lookup_details(rule)
                requested = "; ".join(
                    f"{key}: {display_value(value)}"
                    for key, value in details.items()
                    if value
                )
                description = f"{result['id']}\n{description}\n{requested}\n{result.get('message', '')}"
                # Context is explicitly separate from answer locations and
                # never contributes a cell highlight, comment, or graded link.
                context = next(
                    (
                        item
                        for item in result.get("context_locations", [])
                        if item.get("role") in {"table_header", "section_header"}
                        and item.get("sheet") in workbook.sheetnames
                    ),
                    None,
                )
                if context:
                    description += (
                        f"\nOpen {context['role'].replace('_', ' ')} (not graded): "
                        f"{context['sheet']}!{context['cell']}"
                    )
            for index, location in enumerate(locations.get(result["id"]) or [None]):
                height = max(
                    row_height(description, 57) if index == 0 else 30,
                    row_height(location[0], 28) if location else 44,
                )
                shade = "FFFFFF" if number % 2 else "F7F9FB"
                for column in range(1, 6):
                    text(row, None, column, column, color=shade, height=height)
                    if index == 0:
                        sheet.cell(row, column).border = Border(
                            top=Side(style="hair", color=line)
                        )
                if index == 0:
                    text(
                        row,
                        number,
                        1,
                        1,
                        color=shade,
                        font_color=muted,
                        align="center",
                        height=height,
                    )
                    description_cell = text(
                        row, description, 2, 2, color=shade, height=height
                    )
                    if context:
                        _summary_link(
                            description_cell, context["sheet"], context["cell"]
                        )
                        context_links += 1
                if location:
                    title, coordinate = location
                    text(row, title, 3, 3, color=shade, height=height)
                    cell = text(
                        row,
                        coordinate,
                        4,
                        4,
                        color=tint(color),
                        align="center",
                        height=height,
                    )
                    _summary_link(cell, title, coordinate, bold=True)
                    linked_cells += 1
                elif (
                    rule.get("kind") in UNHIGHLIGHTED_PASSES
                    and result.get("passed") is True
                ):
                    text(row, result["sheet"], 3, 3, color=shade, height=height)
                    text(row, "—", 4, 4, color=shade, font_color=muted, height=height)
                else:
                    text(
                        row,
                        _unlocated_label(rule, result),
                        3,
                        4,
                        color=shade,
                        font_color=muted,
                        height=height,
                    )
                text(
                    row,
                    _verdict_label(rule, result),
                    5,
                    5,
                    color=tint(color),
                    bold=True,
                    align="center",
                    height=height,
                ).border = Border(left=Side(style="medium", color=color))
                row += 1
        sheet.row_dimensions[row].height = 14
        row += 1
    sheet.print_options.horizontalCentered = False
    sheet.print_area = f"A1:E{row - 1}"
    sheet.print_title_rows = "1:2"
    sheet.page_margins = PageMargins(
        left=0.3, right=0.3, top=0.4, bottom=0.4, header=0.15, footer=0.2
    )
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.oddFooter.left.text = "Financial Audit Bench"
    sheet.oddFooter.right.text = "Page &P of &N"
    for footer in (sheet.oddFooter.left, sheet.oddFooter.right):
        footer.size, footer.color = 8, muted
    for tab in workbook:
        tab.sheet_view.tabSelected = tab is sheet
    workbook.active = sheet
    workbook.views[0].firstSheet = 0
    return {
        "sheet": sheet.title,
        "score": score,
        "passed_checks": passed,
        "total_checks": len(results),
        "linked_cells": linked_cells,
        "context_links": context_links,
        "earned_points": totals["earned_points"],
        "possible_points": totals["possible_points"],
    }


def apply_grading_highlights(sheet: Any, colors: dict[str, str]) -> None:
    """Keep grading visible above existing conditional formatting, including merges."""
    existing = []
    for region in list(sheet.conditional_formatting):
        rules = sheet.conditional_formatting[region]
        rules[:] = [rule for rule in rules if rule.formula != [GRADING_FORMULA]]
        if not rules:
            del sheet.conditional_formatting[str(region.sqref)]
        existing.extend(rules)
    areas: dict[str, set[str]] = {}
    for coordinate, color in colors.items():
        sheet[coordinate].fill = PatternFill(fill_type="solid", fgColor="FF" + color)
        area = next(
            (
                str(merged)
                for merged in sheet.merged_cells.ranges
                if coordinate in merged
            ),
            coordinate,
        )
        areas.setdefault(color, set()).add(area)
    # Preserve the relative precedence of original rules below the grading rules.
    for priority, rule in enumerate(
        sorted(existing, key=lambda rule: rule.priority), len(areas) + 1
    ):
        rule.priority = priority
    for priority, (color, ranges) in enumerate(sorted(areas.items()), 1):
        rule = FormulaRule(
            formula=[GRADING_FORMULA],
            stopIfTrue=True,
            fill=PatternFill(fill_type="solid", fgColor="FF" + color),
        )
        rule.priority = priority
        sheet.conditional_formatting.add(" ".join(sorted(ranges)), rule)


def check_type(check: dict, result: dict | None = None) -> str:
    return (
        "llm"
        if (result or {}).get("kind", check.get("kind")) == "llm"
        else "deterministic"
    )


def check_type_label(check: dict, result: dict | None = None) -> str:
    return "LLM" if check_type(check, result) == "llm" else "Deterministic"


def grading_comment(
    existing: Any, explanations: list[str], *, highlight: str = ""
) -> Comment:
    """Replace previous feedback without replacing the auditor's own comment."""
    original = submitted_comment_text(existing) or ""
    feedback = []
    if highlight and len(explanations) > 1:
        feedback.append("Highlight: " + highlight)
    feedback.extend(explanations)
    text = (
        original
        + GRADING_COMMENT_START
        + ("\n\n" if original else "")
        + "\n\n".join(feedback)
        + GRADING_COMMENT_END
    )
    comment = copy(existing) if existing else Comment("", "Financial Audit Bench")
    comment.text = text
    comment.width = max(comment.width, 440)
    lines = sum(max(1, (len(line) + 64) // 65) for line in text.splitlines())
    comment.height = max(70, min(720, 30 + 16 * lines))
    return comment


def check_explanation(check: dict, result: dict, tolerance: Decimal) -> str:
    """Describe the actual rule, including conditional values without JSON dumps."""
    description = (
        check.get("description")
        or result.get("description")
        or "Required workbook value"
    )
    verdict = "Passed" if result["passed"] else "Failed"
    lines = [f"{check_type_label(check, result)} / {verdict}", f"Check: {description}"]
    if result.get("grading_method") == "llm_table":
        lines.append("Method: table review after multiple row matches.")
    if check.get("kind") == "llm":
        expected = check.get("expected", result.get("expected"))
        if expected and expected != description:
            lines.append(f"Expected: {expected}")
    else:
        expected_text = expectation(check, tolerance)
        if expected_text:
            lines.append("Expected: " + expected_text)
    # Passing messages usually repeat the result. Failed explanations are useful.
    if not result["passed"] and result.get("message"):
        lines.append("Reason: " + result["message"])
    if check.get("kind") == "notes_have_context":
        lines.append("Scoring: one point shared across all tables in this workpaper.")
    return "\n".join(lines)


def annotate_workbook(
    workbook_path: Path,
    rubric: dict[str, Any],
    checks: list[dict[str, Any]],
) -> dict[str, Any]:
    """Apply distinct check-type/result highlights above workbook formatting."""
    # Standalone feedback generation may add navigation context, never mutate
    # the persisted grading evidence supplied by the caller.
    checks = [dict(check) for check in checks]
    by_id = {check["id"]: check for check in checks}
    rules = {check["id"]: check for tab in rubric["tabs"] for check in tab["checks"]}
    tolerance = Decimal(
        str(rubric.get("comparison", {}).get("numeric_absolute_tolerance", 0.01))
    )
    workbook = load_workbook(workbook_path)
    applied: list[dict[str, Any]] = []
    try:
        for tab in rubric["tabs"]:
            sheet_name = tab["sheet"]
            for rule in tab["checks"]:
                result = by_id[rule["id"]]
                if type(result.get("passed")) is not bool:
                    continue
                passed = result["passed"]
                if passed and rule.get("kind") in UNHIGHLIGHTED_PASSES:
                    continue
                locations = result.get(
                    "passing_locations" if passed else "matched_locations", []
                )
                by_sheet: dict[str, list[str]] = defaultdict(list)
                for location in locations:
                    if location["sheet"] in workbook.sheetnames:
                        by_sheet[location["sheet"]].append(location["cell"])
                if not passed and not by_sheet:
                    target_title = (
                        rule["extract"][0]["sheet"]
                        if rule.get("kind") == "llm"
                        else sheet_name
                    )
                    if target_title in workbook.sheetnames:
                        try:
                            result["context_locations"] = missing_answer_context(
                                compile_check(rule, target_title),
                                workbook[target_title],
                            )
                        except (ValueError, KeyError):
                            # Navigation is optional; rendering never repairs or evaluates a rule.
                            result["context_locations"] = []
                for title, coordinates in by_sheet.items():
                    applied.append(
                        {
                            "check_id": rule["id"],
                            "passed": passed,
                            "sheet": title,
                            "cells": list(dict.fromkeys(coordinates)),
                        }
                    )
        # One note per cell contains every contributing check, including failures
        # sharing a target. Merged ranges store notes only at their anchor.
        comments: dict[tuple[str, str], dict[str, str]] = defaultdict(dict)
        for annotation in applied:
            title = annotation["sheet"]
            check_id = annotation["check_id"]
            annotation["check_type"] = check_type(rules[check_id], by_id[check_id])
            for coordinate in annotation["cells"]:
                cell = workbook[title][coordinate]
                for merged in workbook[title].merged_cells.ranges:
                    if coordinate in merged:
                        cell = workbook[title].cell(merged.min_row, merged.min_col)
                        break
                comments[(title, cell.coordinate)][check_id] = check_explanation(
                    rules[check_id], by_id[check_id], tolerance
                )
        colors_by_sheet: dict[str, dict[str, str]] = defaultdict(dict)
        for (title, coordinate), explanations in comments.items():
            cell = workbook[title][coordinate]
            # Failures win, then deterministic checks; independent of rubric order.
            dominant = min(
                explanations,
                key=lambda key: (
                    by_id[key]["passed"],
                    check_type(rules[key], by_id[key]) == "llm",
                ),
            )
            outcome = "passed" if by_id[dominant]["passed"] else "failed"
            color = CHECK_COLORS[check_type(rules[dominant], by_id[dominant])][outcome]
            colors_by_sheet[title][coordinate] = color
            cell.comment = grading_comment(
                cell.comment,
                list(explanations.values()),
                highlight=f"{check_type_label(rules[dominant], by_id[dominant])} / {outcome.title()}",
            )
        for sheet in workbook:
            apply_grading_highlights(sheet, colors_by_sheet[sheet.title])
        summary = add_grading_summary(workbook, rubric, checks, comments)
        for sheet in workbook:
            for view in sheet.views.sheetView:
                view.pane = None
                view.selection = [Selection()]
        workbook.save(workbook_path)
    finally:
        workbook.close()

    return {
        "applied": True,
        "workbook": rubric["output"],
        "colors": CHECK_COLORS,
        "summary": summary,
        "grading_overrides_conditional_formatting": True,
        "shared_cell_rule": SHARED_CELL_RULE,
        "checks_annotated": len({item["check_id"] for item in applied}),
        "cells_annotated": sum(len(item["cells"]) for item in applied),
        "cells_commented": len(comments),
        "details": applied,
    }


def export_feedback(report: dict, source: Path, rubric: dict, output: Path) -> None:
    """Preserve completed grading even when its optional presentation fails."""
    report.pop("annotated_workbook", None)
    report["workbook_highlighting"] = {"applied": False, "status": "pending"}
    write_json(output, report)
    try:
        annotated = output.parent / "annotated" / rubric["output"]
        annotated.parent.mkdir(parents=True, exist_ok=True)
        with tempfile.TemporaryDirectory(
            prefix=".feedback-", dir=annotated.parent
        ) as directory:
            temporary = Path(directory) / rubric["output"]
            shutil.copyfile(source, temporary)
            details = annotate_workbook(temporary, rubric, report["checks"])
            temporary.replace(annotated)
        report["workbook_highlighting"] = {**details, "status": "complete"}
        report["annotated_workbook"] = str(annotated)
    except Exception as error:
        report["workbook_highlighting"] = {
            "applied": False,
            "status": "error",
            "error": {"type": type(error).__name__, "message": str(error)},
        }
    write_json(output, report)
